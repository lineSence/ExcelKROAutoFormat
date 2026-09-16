"""Файл продавцов из 1С: пары «ФИО — часы», итог, склад, русские ошибки.

Данные в тестах синтетические и обезличенные (`ФИО 1` и так далее),
реальные выгрузки в репозиторий не кладём.
"""

import pytest
from openpyxl import Workbook

from app.core import sellers
from app.core.meta import SHARE_HOURS, SheetMeta
from app.services.upload_job import UPLOAD_STAGES

PAIRS = (("ФИО 1", 12), ("ФИО 2", 216), ("ФИО 3", 24))


def _sheet(pairs=PAIRS, total=True, skip_hours=()):
    """Лист в том же виде, в каком его отдаёт 1С."""
    book = Workbook()
    sheet = book.active
    sheet.title = "TDSheet"
    sheet["A2"] = "Параметры:"
    sheet["C2"] = "Дата кон: 15.09.2026 23:59:59"
    sheet["C3"] = "Дата нач: 11.08.2026 0:00:00"
    sheet["A4"] = "Отбор:"
    sheet["C4"] = 'Склад Равно "ТестовыйСкладПДВ"'
    sheet["A6"] = "Продавец"
    sheet["A7"] = "Часы"
    row = 9
    for name, hours in pairs:
        sheet.cell(row=row, column=1, value=name)
        if name not in skip_hours:
            sheet.cell(row=row + 1, column=1, value=hours)
        row += 2
    if total:
        sheet.cell(row=row, column=1, value=sum(item[1] for item in pairs))
    return sheet


def test_sellers_pairs_and_warehouse():
    data = sellers.read_sheet(_sheet())

    assert [(item.name, item.hours) for item in data.sellers] == [
        ("ФИО 1", "12"),
        ("ФИО 2", "216"),
        ("ФИО 3", "24"),
    ]
    assert data.warehouse == "ТестовыйСкладПДВ"
    assert data.hours_sum == 252
    assert data.notes == []


def test_total_row_is_not_a_seller():
    """Итоговая строка без ФИО в список не попадает, но сохраняется как итог."""
    data = sellers.read_sheet(_sheet())

    assert len(data.sellers) == 3
    assert data.total == 252
    assert data.total == data.hours_sum


def test_headers_and_filters_are_skipped():
    data = sellers.read_sheet(_sheet())
    names = [item.name for item in data.sellers]

    assert "Продавец" not in names
    assert "Часы" not in names
    assert not any(name.startswith("Отбор") or name.startswith("Склад") for name in names)


def test_missing_hours_leave_russian_note():
    data = sellers.read_sheet(_sheet(total=False, skip_hours=("ФИО 2",)))

    assert ("ФИО 2", "") in [(item.name, item.hours) for item in data.sellers]
    assert data.notes
    assert "ФИО 2" in data.notes[0]
    assert "вручную" in data.notes[0]


def test_note_without_sellers():
    book = Workbook()
    sheet = book.active
    sheet["A1"] = "Параметры:"

    data = sellers.read_sheet(sheet)

    assert data.sellers == ()
    assert data.note() == "продавцы в файле не найдены"


def test_warehouse_from_filter_variants():
    assert sellers.warehouse_from_filter('Склад Равно "АБВ"') == "АБВ"
    assert sellers.warehouse_from_filter("Склад Равно «АБВ»") == "АБВ"
    assert sellers.warehouse_from_filter("Дата нач: 11.08.2026") == ""


def test_hours_with_comma_and_spaces():
    sheet = _sheet(pairs=(("ФИО 1", 0),), total=False)
    sheet.cell(row=10, column=1, value="12,5")

    data = sellers.read_sheet(sheet)

    assert data.sellers[0].hours == "12.5"


def test_broken_file_gives_russian_error(tmp_path):
    """Не архив — русская ошибка вместо трассировки."""
    broken = tmp_path / "sellers.xlsx"
    broken.write_bytes(b"not a zip")

    with pytest.raises(sellers.SellersError) as error:
        sellers.read(broken, tmp_path / "work")

    assert "Файл продавцов не прочитан" in str(error.value)


def test_upload_has_sellers_stage():
    assert "sellers" in dict(UPLOAD_STAGES)


def test_sheet_meta_counts_shares_by_hours():
    data = sellers.read_sheet(_sheet())
    meta = SheetMeta(sellers=tuple(data.sellers), share_mode=SHARE_HOURS)

    assert meta.share_mode == SHARE_HOURS
    assert len(meta.sellers) == 3
