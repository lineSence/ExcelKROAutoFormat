"""Пометка «нет фото» в готовом файле сверки."""

from __future__ import annotations

import openpyxl

from app.core import photo_mark

SHEET = "TDSheet"
# Строки данных: плюс, минус, плюс.
ROWS = (
    (5, "К-1", "Жидкость А", 2),
    (6, "К-2", "Жидкость Б", -2),
    (7, "К-3", "Жидкость В", 1),
)


def _make_file(path):
    """Мини-сверка: титул, заголовок группы, три строки и итог."""
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = SHEET
    sheet.cell(row=1, column=1).value = "Инвентаризация товаров № ИНВ-15 от 09.09.2026"
    sheet.cell(row=4, column=1).value = "№"
    sheet.cell(row=4, column=2).value = "Жидкости"
    sheet.cell(row=4, column=3).value = "Хар-ка"
    for row, code, name, diff in ROWS:
        sheet.cell(row=row, column=1).value = code
        sheet.cell(row=row, column=2).value = name
        sheet.cell(row=row, column=4).value = 10
        sheet.cell(row=row, column=5).value = 10 - diff
        sheet.cell(row=row, column=6).value = diff
        sheet.cell(row=row, column=10).value = 0
    # Строка итога: код и наименование пусты, есть число.
    sheet.cell(row=8, column=4).value = 30
    book.save(path)
    book.close()
    return path


def _notes(path) -> dict[int, str]:
    book = openpyxl.load_workbook(path)
    sheet = book[SHEET]
    values = {
        row: str(sheet.cell(row=row, column=photo_mark.COLUMN).value or "")
        for row, *_ in ROWS
    }
    book.close()
    return values


def test_plus_rows_get_note(tmp_path):
    """Пометка появляется только у плюсующих строк."""
    file = _make_file(tmp_path / "сверка.xlsx")

    ok, detail = photo_mark.apply(file, sheet_name=SHEET)

    assert ok, detail
    notes = _notes(file)
    assert notes[5] == "нет фото"
    assert notes[7] == "нет фото"
    assert notes[6] == ""


def test_confirmed_row_loses_note(tmp_path):
    """Подтверждённое фото снимает пометку только со своей строки."""
    file = _make_file(tmp_path / "сверка.xlsx")
    photo_mark.apply(file, sheet_name=SHEET)

    ok, _ = photo_mark.apply(file, confirmed=[5], sheet_name=SHEET)

    assert ok
    notes = _notes(file)
    assert notes[5] == ""
    assert notes[7] == "нет фото"


def test_note_joins_existing_comment_once(tmp_path):
    """Пометка дописывается к комментарию реестра и не дублируется."""
    file = _make_file(tmp_path / "сверка.xlsx")
    book = openpyxl.load_workbook(file)
    book[SHEET].cell(row=5, column=photo_mark.COLUMN).value = "недовоз ПДВ06030196 3шт."
    book.save(file)
    book.close()

    photo_mark.apply(file, sheet_name=SHEET)
    photo_mark.apply(file, sheet_name=SHEET)

    assert _notes(file)[5] == "недовоз ПДВ06030196 3шт., нет фото"


def test_old_note_is_replaced(tmp_path):
    """Пометка прежних версий «есть фото» снимается."""
    file = _make_file(tmp_path / "сверка.xlsx")
    book = openpyxl.load_workbook(file)
    book[SHEET].cell(row=7, column=photo_mark.COLUMN).value = "есть фото"
    book.save(file)
    book.close()

    photo_mark.apply(file, sheet_name=SHEET)

    assert _notes(file)[7] == "нет фото"


def test_missing_file_is_reported(tmp_path):
    """Удалённый файл — понятное сообщение, а не исключение."""
    ok, detail = photo_mark.apply(tmp_path / "нет.xlsx", sheet_name=SHEET)

    assert not ok
    assert "удал" in detail
