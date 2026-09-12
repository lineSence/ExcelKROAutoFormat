"""Тесты ядра. Файл сверки собирается в коде теста."""

from __future__ import annotations

import os
import re
import sys
import time
import zipfile
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.config import Settings
from app.core.format import output_filename
from app.core.parse import parse, warehouse_from_filename
from app.core.pipeline import process, sweep
from app.core.repair import describe, needs_repair, repair_by_inject
from app.core.resort import normalize

GROUPS = {
    "Сигареты": [
        ("Сигареты Parliament Aqua Blue", -3, -450.0),
        ("Сигареты Parliament Aqua Blue КС", 3, 450.0),
        ("Сигареты Winston XS Blue", -2, -300.0),
        ("Сигареты Kent Nano White", 4, 600.0),
    ],
    "Стики": [
        ("Стики Heets Amber", -1, -120.0),
        ("Стики HEETS Amber Selection", 1, 120.0),
    ],
}


def _build_source(path: Path) -> None:
    """Собирает файл, похожий на выгрузку 1С."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "TDSheet"

    sheet["A1"] = "Инвентаризация товаров № ИНВ-15 от 09.09.2026"
    sheet["A2"] = "Организация: ООО Торг"
    sheet["A3"] = "Склад:"
    sheet["H3"] = "Валюта"
    sheet["I3"] = "руб"

    row = 5
    for name, items in GROUPS.items():
        sheet.cell(row=row, column=1).value = "№"
        sheet.cell(row=row, column=2).value = name
        sheet.cell(row=row, column=3).value = "Хар-ка"
        sheet.merge_cells(start_row=row, start_column=2, end_row=row + 1, end_column=2)
        row += 2
        for number, (item_name, diff, sum_diff) in enumerate(items, start=1):
            sheet.cell(row=row, column=1).value = number
            sheet.cell(row=row, column=2).value = item_name
            sheet.cell(row=row, column=4).value = 10
            sheet.cell(row=row, column=5).value = 10 - diff
            sheet.cell(row=row, column=6).value = diff
            sheet.cell(row=row, column=7).value = 150.0
            sheet.cell(row=row, column=10).value = sum_diff
            sheet.cell(row=row, column=11).value = f"Документ {number}"
            row += 1
        sheet.cell(row=row, column=4).value = len(items) * 10
        sheet.cell(row=row, column=6).value = sum(item[1] for item in items)
        sheet.cell(row=row, column=10).value = sum(item[2] for item in items)
        row += 2

    workbook.save(path)


def _rewrite(path: Path, change) -> None:
    with zipfile.ZipFile(path) as archive:
        items = [(item, archive.read(item.filename)) for item in archive.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for item, data in items:
            new_data = change(item.filename, data)
            if new_data is None:
                continue
            archive.writestr(item, new_data)


def _drop_shared_strings(path: Path) -> None:
    """Убирает xl/sharedStrings.xml, как в выгрузке 1С."""
    _rewrite(path, lambda name, data: None if name == "xl/sharedStrings.xml" else data)


def _break_cell_styles(path: Path) -> None:
    """Стили ячеек ссылаются за пределы cellXfs — болезнь выгрузки 1С."""

    def change(name: str, data: bytes) -> bytes:
        if name.startswith("xl/worksheets/"):
            return re.sub(rb'<c r="B(\d+)"', rb'<c s="99" r="B\1"', data)
        return data

    _rewrite(path, change)


def _break_fonts(path: Path) -> None:
    """Записи стилей ссылаются на несуществующий шрифт."""

    def change(name: str, data: bytes) -> bytes:
        if name == "xl/styles.xml":
            return re.sub(rb'fontId="\d+"', b'fontId="77"', data)
        return data

    _rewrite(path, change)


@pytest.fixture()
def source_file(tmp_path: Path) -> Path:
    path = tmp_path / "ОхтаМоллСМА без форматирования.xlsx"
    _build_source(path)
    _drop_shared_strings(path)
    _break_cell_styles(path)
    return path


def test_warehouse_from_filename() -> None:
    assert warehouse_from_filename("ОхтаМоллСМА без форматирования.xlsx") == "ОхтаМоллСМА"
    assert warehouse_from_filename("БалтийскийТЦМДБ.xlsx") == "БалтийскийТЦМДБ"
    # Имя из нескольких слов берётся целиком.
    assert warehouse_from_filename("Красные Ворота.xlsx") == "Красные Ворота"
    assert warehouse_from_filename("Красные Ворота 09.09.2026 (1).xlsx") == "Красные Ворота"


def test_output_filename() -> None:
    assert output_filename("ОхтаМоллСМА", "09.09.2026") == "ОхтаМоллСМА 09.09.2026.xlsx"


def test_normalize_folds_names() -> None:
    words = ("сигареты", "стики")
    first = normalize("Сигареты Parliament Aqua Blue", words)
    second = normalize("Сигареты  Parliament Aqua Blue КС", words)
    assert first
    assert second.startswith(first)


def test_broken_file_fails_without_repair(source_file: Path) -> None:
    """Без ремонта openpyxl файл не открывает.

    Тип ошибки зависит от версии openpyxl, поэтому проверяется любая из трёх.
    """
    with pytest.raises((IndexError, KeyError, ValueError)):
        openpyxl.load_workbook(source_file)


def test_repair_fixes_styles_and_strings(source_file: Path, tmp_path: Path) -> None:
    assert needs_repair(source_file) is True
    assert describe(source_file)["style_problems"] > 0

    repaired = repair_by_inject(source_file, tmp_path / "repaired.xlsx")
    assert needs_repair(repaired) is False

    sheet = openpyxl.load_workbook(repaired)["TDSheet"]
    assert sheet["A1"].value.startswith("Инвентаризация товаров")


def test_repair_fixes_broken_fonts(tmp_path: Path) -> None:
    path = tmp_path / "ОхтаМоллСМА без форматирования.xlsx"
    _build_source(path)
    _break_fonts(path)

    assert needs_repair(path) is True
    repaired = repair_by_inject(path, tmp_path / "repaired-fonts.xlsx")
    openpyxl.load_workbook(repaired)


def test_parse_finds_groups(source_file: Path, tmp_path: Path) -> None:
    repaired = repair_by_inject(source_file, tmp_path / "r.xlsx")
    sheet = openpyxl.load_workbook(repaired)["TDSheet"]
    document = parse(sheet)
    assert document.doc_date == "09.09.2026"
    assert len(document.groups) == len(GROUPS)


def test_process_makes_output(source_file: Path, tmp_path: Path) -> None:
    settings = Settings.load()
    settings.tmp_dir = str(tmp_path / "work")
    result = process(source_file, source_file.name, settings)

    assert result.output_name == "ОхтаМоллСМА 09.09.2026.xlsx"
    assert result.output_path.is_file()
    assert result.summary["groups"] == len(GROUPS)
    assert result.summary["pieces"] > 0

    sheet = openpyxl.load_workbook(result.output_path)["TDSheet"]
    assert str(sheet["I4"].value).startswith("=SUM(")
    assert sheet["C4"].value == "ОхтаМоллСМА"
    assert sheet.auto_filter.ref.startswith("K1:K")


def test_process_uses_given_folder(source_file: Path, tmp_path: Path) -> None:
    """Веб-слой даёт свою папку: лишние папки не создаются."""
    settings = Settings.load()
    settings.tmp_dir = str(tmp_path / "work")
    folder = Path(settings.tmp_dir) / "one"
    result = process(source_file, source_file.name, settings, folder=folder)

    assert result.output_path.parent == folder
    assert [item.name for item in Path(settings.tmp_dir).iterdir()] == ["one"]


def test_sweep_removes_old_folders(tmp_path: Path) -> None:
    settings = Settings.load()
    settings.tmp_dir = str(tmp_path / "work")
    settings.result_ttl_minutes = 1

    old = Path(settings.tmp_dir) / "old"
    fresh = Path(settings.tmp_dir) / "fresh"
    old.mkdir(parents=True)
    fresh.mkdir(parents=True)
    past = time.time() - 3600
    os.utime(old, (past, past))

    assert sweep(settings) == 1
    assert not old.exists()
    assert fresh.exists()
