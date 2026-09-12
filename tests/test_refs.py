"""Тесты справочников: чтение книг, поиск пары и ручная загрузка."""

from __future__ import annotations

import datetime as dt

import openpyxl

from app.core import refs, refs_sync


def make_planning(path):
    """Книга планирования: лист администратора, строки — магазины."""
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Казаков А."
    sheet["A1"] = 2026
    sheet["A2"] = "№"
    sheet["B2"] = "имя"
    sheet["C2"] = "Последний переучет"
    sheet["D2"] = "Август 10-16"
    sheet["E2"] = "Август 17-23"
    sheet["A3"] = 1
    sheet["B3"] = "ОхтаМоллСМА (НОЧЬ)"
    sheet["D3"] = "планово"
    sheet["E3"] = "увольнение"
    guide = book.create_sheet("обозначения")
    guide["A1"] = "Планово"
    book.save(path)


def make_schedule(path):
    """Книга графика: лист распределения и лист полных ФИО."""
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = refs.SCHEDULE_SHEET
    sheet["A1"] = "дата"
    sheet["B1"] = "Казаков А."
    sheet["C1"] = "Савин А."
    sheet["D1"] = "Козлова А."
    sheet["A2"] = dt.datetime(2026, 8, 12)
    sheet["B2"] = "ОхтаМоллСМА\nСмирнова\nПавлова М."
    sheet["A3"] = dt.datetime(2026, 8, 13)
    sheet["C3"] = "БалтийскийТЦМДБ\nТарасова"

    phones = book.create_sheet("тлф")
    phones["A1"] = "Смирнова Ольга Ивановна"
    phones["A2"] = "Павлова Марина Владимировна"
    phones["A3"] = "Павлова Екатерина Александровна"
    book.save(path)


def books(tmp_path) -> refs.RefsBooks:
    planning = tmp_path / "planning.xlsx"
    schedule = tmp_path / "schedule.xlsx"
    make_planning(planning)
    make_schedule(schedule)
    refs.forget_books()
    return refs.load_books(str(planning), str(schedule), force=True)


def test_normalize_store_drops_notes():
    assert refs.normalize_store("ОхтаМоллСМА (НОЧЬ)") == refs.normalize_store("ОхтаМоллСМА")


def test_reason_and_auditors_found(tmp_path):
    data = books(tmp_path)
    info = refs.lookup("ОхтаМоллСМА", dt.date(2026, 8, 12), data, checker="Разумовский")
    assert info.reason == "планово"
    assert info.admin == "Казаков А."
    assert info.checker == "Разумовский"
    # Фамилия без инициала и фамилия с инициалом развёрнуты в полные ФИО.
    assert info.auditors == (
        "Смирнова Ольга Ивановна",
        "Павлова Марина Владимировна",
    )


def test_unknown_pair_stays_empty(tmp_path):
    data = books(tmp_path)
    info = refs.lookup("НетТакогоСклада", dt.date(2026, 8, 12), data, checker="Разумовский")
    assert info.reason == ""
    assert info.admin == ""
    assert info.auditors == ()
    assert info.found is False


def test_write_cells_skips_empty(tmp_path):
    book = openpyxl.Workbook()
    sheet = book.active
    info = refs.RefsInfo(reason="", admin="", checker="Разумовский", auditors=())
    written = refs.write_cells(
        sheet,
        info,
        {"reason": "C5", "admin": "L2", "checker": "L3", "auditors": "L4"},
    )
    # Пустые значения в файл не попадают и никаких пометок не появляется.
    assert written == ["L3"]
    assert sheet["C5"].value is None
    assert sheet["L2"].value is None
    assert sheet["L4"].value is None


def test_full_name_keeps_ambiguous():
    index = {
        "павлова": ["Павлова Марина Владимировна", "Павлова Екатерина Александровна"],
    }
    assert refs.full_name("Павлова", index) == "Павлова"
    assert refs.full_name("Павлова Е.", index) == "Павлова Екатерина Александровна"


def test_state_round_trip(tmp_path):
    path = tmp_path / "state.json"
    state = refs_sync.SyncState(
        planning_name="НОВ ПЛАНИРОВАНИЕ.xlsx",
        days=[0, 2],
        times=["7:5", "19:00"],
    )
    refs_sync.save_state(path, state)
    back = refs_sync.load_state(path)
    assert back.planning_name == "НОВ ПЛАНИРОВАНИЕ.xlsx"
    assert back.days == [0, 2]
    assert back.times == ["07:05", "19:00"]


def test_book_status_sees_loaded_files(tmp_path):
    make_planning(refs_sync.book_target(tmp_path, "planning"))
    report = refs_sync.book_status(tmp_path)
    assert report["planning"]["found"] is True
    assert report["schedule"]["found"] is False


def test_mark_upload_saves_state(tmp_path):
    state_path = tmp_path / "state.json"
    state = refs_sync.SyncState()
    result = refs_sync.mark_upload(state, "schedule", "График.xlsx", state_path)
    assert result.schedule_name == "График.xlsx"
    assert result.schedule_loaded
    assert result.last_status == "загружено вручную"
    assert refs_sync.load_state(state_path).schedule_name == "График.xlsx"


def test_due_respects_days_and_times():
    # Расписание пока не используется, но правила должны работать.
    state = refs_sync.SyncState(days=[2], times=["07:30"], enabled=True)
    # Среда после 07:30 — копия нужна.
    assert refs_sync.due(state, dt.datetime(2026, 9, 9, 7, 31), None) is True
    # Та же среда, копия уже была.
    assert refs_sync.due(state, dt.datetime(2026, 9, 9, 7, 31), dt.datetime(2026, 9, 9, 7, 30, 5)) is False
    # Другой день недели.
    assert refs_sync.due(state, dt.datetime(2026, 9, 10, 9, 0), None) is False
    # Расписание выключено.
    state.enabled = False
    assert refs_sync.due(state, dt.datetime(2026, 9, 9, 7, 31), None) is False
