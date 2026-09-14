"""Пометка «нет фото» в готовом файле сверки.

Отметка нужна не на том товаре, по которому фото пришло, а на том, по
которому фото нет: в столбец 12 (`L`) каждой плюсующей строки (`F > 0`) без
подтверждённого снимка пишется «нет фото». Подтверждение снимка убирает
пометку с этой строки.

Пометки ставятся после каждой сборки файла, потому что пересборка от
исходника (подтверждение спорных пар, ручные поля, заявки реестра) пишет
столбец 12 заново. Поэтому модуль умеет и ставить, и снимать пометку, а
список строк с подтверждённым фото живёт в памяти веб-слоя по токену сверки.

Модуль правит уже собранный файл, поэтому ошибки не поднимаются наверх:
возвращается пара «получилось, пояснение». Ответ человека важнее пометки:
если файл не открылся, пример всё равно запишется.
"""

from __future__ import annotations

import logging
from collections.abc import Iterable
from pathlib import Path

import openpyxl

from .parse import COL_DIFF, ParseError, parse

logger = logging.getLogger("excelkro.photo_mark")

# Столбец 12 — `L`, служебный комментарий сверки.
COLUMN = 12
NOTE = "нет фото"
# Пометка прежних версий: снимается, чтобы в файле не осталось двух ответов.
OLD_NOTE = "есть фото"


def _sheet(workbook, sheet_name: str):
    if sheet_name and sheet_name in workbook.sheetnames:
        return workbook[sheet_name]
    return workbook[workbook.sheetnames[0]]


def _is_number(value: object) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def plus_rows(sheet) -> list[int]:
    """Строки данных, где итог в столбце `F` больше нуля.

    Строки берутся из готового файла, а не из подсказок распознавания:
    пометка нужна всем плюсующим позициям, даже если фото не присылали.
    """
    try:
        document = parse(sheet)
    except ParseError as error:
        logger.warning("Плюсующие строки не найдены: %s", error)
        return []
    rows: list[int] = []
    for group in document.groups:
        for row in group.data_rows:
            value = sheet.cell(row=row, column=COL_DIFF).value
            if _is_number(value) and value > 0:
                rows.append(row)
    return rows


def _parts(cell) -> list[str]:
    parts = [part.strip() for part in str(cell.value or "").split(",")]
    return [part for part in parts if part]


def _write(cell, parts: list[str]) -> None:
    cell.value = ", ".join(parts) if parts else None


def _add(cell, note: str) -> bool:
    """Дописывает пометку через запятую. Повтор ничего не дублирует."""
    parts = _parts(cell)
    if any(part.casefold() == note.casefold() for part in parts):
        return False
    parts.append(note)
    _write(cell, parts)
    return True


def _remove(cell, notes: Iterable[str]) -> bool:
    """Снимает перечисленные пометки, остальной текст ячейки сохраняется."""
    drop = {str(note).casefold() for note in notes}
    parts = _parts(cell)
    kept = [part for part in parts if part.casefold() not in drop]
    if len(kept) == len(parts):
        return False
    _write(cell, kept)
    return True


def apply(
    path: str | Path,
    confirmed: Iterable[int] = (),
    sheet_name: str = "",
    note: str = NOTE,
    column: int = COLUMN,
) -> tuple[bool, str]:
    """Ставит «нет фото» плюсующим строкам без подтверждённого снимка."""
    file = Path(path)
    if not file.is_file():
        return False, "Файл сверки уже удалён: пометки о фото не записаны."

    have = {int(row) for row in confirmed if int(row or 0) > 0}

    try:
        workbook = openpyxl.load_workbook(file)
    except Exception as error:  # noqa: BLE001
        logger.warning("Файл сверки не открыт: %s", error)
        return False, f"Файл сверки не открылся: {error}"

    marked = 0
    try:
        sheet = _sheet(workbook, sheet_name)
        rows = plus_rows(sheet)
        for row in rows:
            cell = sheet.cell(row=row, column=column)
            if row in have:
                _remove(cell, (note, OLD_NOTE))
                continue
            _remove(cell, (OLD_NOTE,))
            _add(cell, note)
            marked += 1
        workbook.save(file)
    except Exception as error:  # noqa: BLE001
        logger.exception("Пометки о фото не записаны")
        return False, f"Пометки о фото не записаны: {type(error).__name__}: {error}"
    finally:
        workbook.close()

    kept = len(have & set(rows))
    return True, (
        f"Пометка «{note}» стоит в строках: {marked}. "
        f"Строк с подтверждённым фото: {kept}."
    )
