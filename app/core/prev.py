"""Чтение предыдущей инвентаризации для сравнения с текущей сверкой.

Предыдущая сверка — это готовый (уже разобранный) файл прошлой
инвентаризации. Из него нужны три вещи:

* позиции: название, количество расхождения, сумма и пометка в столбце C
  («не+», «перез»);
* признак «считалась в минус»: количество меньше нуля, а сумма не обнулена,
  то есть эту сумму списывали с продавцов;
* блок продавцов внизу файла: ФИО и вес (часы или единица), по которым
  сумма делилась в прошлый раз.

Позиции сопоставляются по полному названию товара (столбец B).
"""

from __future__ import annotations

import logging
import re
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

from .parse import (
    COL_BOOK,
    COL_CODE,
    COL_DIFF,
    COL_FACT,
    COL_NAME,
    COL_PRICE,
    COL_SUM_DIFF,
    COL_TRAIT,
    parse,
)

logger = logging.getLogger("excelkro.prev")

MARK_NOT_PLUS = "не+"
SPACES = re.compile(r"\s+")

# Подписи и заголовки мини-таблиц: это не продавцы.
SERVICE_LABELS = (
    "проверил",
    "администратор",
    "ревизоры",
    "ночной продавец",
    "перезачёт",
    "перезачет",
    "итого",
)
# Блок продавцов идёт сразу под итогом последней группы.
SELLER_SCAN_ROWS = 80
# Числа позиции: у строки продавца эти столбцы пустые.
ITEM_COLUMNS = (COL_FACT, COL_BOOK, COL_DIFF, COL_PRICE)


def name_key(value: object) -> str:
    """Ключ позиции: полное название без регистра и лишних пробелов."""
    text = str(value or "").replace("\xa0", " ").strip()
    return SPACES.sub(" ", text).lower()


def _number(value: object) -> float:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    return 0.0


def _is_number(value: object) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


@dataclass(frozen=True)
class PrevItem:
    """Позиция предыдущей инвентаризации."""

    name: str
    diff: float
    sum_diff: float
    mark: str

    @property
    def was_minus(self) -> bool:
        """Минус с неснятой суммой: её списывали с продавцов."""
        return self.diff < 0 and round(self.sum_diff, 2) != 0

    @property
    def was_not_plus(self) -> bool:
        """Позиция была помечена «не+»."""
        return MARK_NOT_PLUS in self.mark


@dataclass(frozen=True)
class PrevSeller:
    """Продавец предыдущей сверки и его вес при списании."""

    name: str
    weight: float = 0.0


@dataclass
class PrevSheet:
    """Разобранная предыдущая инвентаризация."""

    file_name: str = ""
    date: str = ""
    items: dict[str, PrevItem] = field(default_factory=dict)
    sellers: tuple[PrevSeller, ...] = ()

    @property
    def total_weight(self) -> float:
        return sum(seller.weight for seller in self.sellers)


def _merge(first: PrevItem, second: PrevItem) -> PrevItem:
    """Одно и то же название встретилось дважды: складываем."""
    marks = " ".join(part for part in (first.mark, second.mark) if part)
    return PrevItem(
        name=first.name,
        diff=first.diff + second.diff,
        sum_diff=first.sum_diff + second.sum_diff,
        mark=marks,
    )


def _is_formula(value: object) -> bool:
    return isinstance(value, str) and value.startswith("=")


def _looks_like_item(sheet, row: int) -> bool:
    """Строка товара: есть код в столбце A или числа в количествах и цене."""
    if str(sheet.cell(row=row, column=COL_CODE).value or "").strip():
        return True
    return any(
        _is_number(sheet.cell(row=row, column=column).value)
        for column in ITEM_COLUMNS
    )


def _is_service_label(sheet, row: int) -> bool:
    return name_key(sheet.cell(row=row, column=COL_NAME).value) in SERVICE_LABELS


def _seller_at(sheet, row: int) -> PrevSeller | None:
    """Продавец: ФИО без чисел позиции, а под ним числовая строка веса."""
    value = sheet.cell(row=row, column=COL_NAME).value
    text = str(value or "").strip()
    if not text or _is_number(value) or _is_formula(value):
        return None
    if _is_service_label(sheet, row) or _looks_like_item(sheet, row):
        return None
    below = sheet.cell(row=row + 1, column=COL_NAME).value
    if not _is_number(below):
        return None
    return PrevSeller(name=text, weight=_number(below))


def read_sellers(sheet, start_row: int) -> tuple[PrevSeller, ...]:
    """Блок продавцов: строка ФИО, под ней строка веса.

    Товары нижней группы в список не попадают, даже если у группы нет
    строки итога. Мини-таблица «Перезачёт» идёт ниже блока продавцов и
    обрывает поиск: её ФИО — это уже другой расчёт.
    """
    result: list[PrevSeller] = []
    row = start_row
    limit = min(sheet.max_row, start_row + SELLER_SCAN_ROWS)
    while row <= limit:
        seller = _seller_at(sheet, row)
        if seller is not None:
            result.append(seller)
            row += 2
            continue
        if result and _is_service_label(sheet, row):
            break
        row += 1
    return tuple(result)


def read(path: str | Path, file_name: str = "", sheet_name: str = "TDSheet") -> PrevSheet:
    """Читает предыдущую сверку. Ошибки разбора пробрасываются наверх."""
    workbook = openpyxl.load_workbook(path, data_only=True)
    if sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
    else:
        sheet = workbook[workbook.sheetnames[0]]

    document = parse(sheet)
    items: dict[str, PrevItem] = {}
    for group in document.groups:
        for row in group.data_rows:
            raw = sheet.cell(row=row, column=COL_NAME).value
            key = name_key(raw)
            if not key:
                continue
            item = PrevItem(
                name=str(raw or "").strip(),
                diff=_number(sheet.cell(row=row, column=COL_DIFF).value),
                sum_diff=_number(sheet.cell(row=row, column=COL_SUM_DIFF).value),
                mark=str(sheet.cell(row=row, column=COL_TRAIT).value or "").strip().lower(),
            )
            found = items.get(key)
            items[key] = _merge(found, item) if found is not None else item

    last_row = max(
        max(group.total_row, group.last_data_row) for group in document.groups
    )
    sellers = read_sellers(sheet, last_row)
    logger.info(
        "Предыдущая сверка: позиций %s, продавцов %s", len(items), len(sellers)
    )
    return PrevSheet(
        file_name=file_name or Path(path).name,
        date=document.doc_date or "",
        items=items,
        sellers=sellers,
    )
