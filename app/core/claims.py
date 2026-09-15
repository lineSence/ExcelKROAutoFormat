"""Реестр расхождений: недовозы и перетарки по текущему магазину.

Реестр — книга Excel «Реестр расхождений» (лист «Исправления»), где склад
и операторы ведут заявки: чего не довезли (недовоз) и что привезли лишним
(перетарка). Такие позиции не вина магазина, поэтому в сверке они:

* получают пометку «не-» (недовоз) или «не+» (перетарка) в столбце C;
* получают разницу суммы 0 в столбце J;
* выделяются зелёным по всей ширине строки (A:N);
* получают комментарий в столбце N вида «недовоз ПДВ06030196 3шт.»;
* не участвуют в подборе пересортов.

Правила разбора согласованы с заказчиком:

* берётся только текущий магазин (имя склада из имени файла сверки);
* окно дат — один год строго назад от даты инвентаризации;
* учитываются только строки с непустым количеством недовоза или перетарки;
* название в заявке должно совпадать с товаром сверки после нормализации
  один в один, и подходить должна ровно одна строка сверки;
* количество по одинаковым названиям суммируется;
* ничего не пишется в файл без подтверждения человека — по каждой строке
  отдельно (страница результата, форма «Заявки реестра»).

Чтение книги тяжёлое (реестр — это тысячи строк), поэтому разобранные
строки складываются в JSON-кеш рядом с книгой и перечитываются только
при смене размера или времени правки файла.
"""

from __future__ import annotations

import datetime as dt
import json
import logging
import threading
from dataclasses import dataclass, field
from hashlib import md5
from pathlib import Path

import openpyxl

from . import refs, resort, style
from .parse import COL_DIFF, COL_NAME, COL_SUM_DIFF, COL_TRAIT

logger = logging.getLogger("excelkro.claims")

# Книга реестра и кеш разобранных строк лежат в папке справочников.
CLAIMS_FILE = "claims.xlsx"
CACHE_FILE = "claims-index.json"
SHEET = "Исправления"

# Окно дат: один год строго назад от даты инвентаризации.
WINDOW_DAYS = 365

KIND_SHORT = "недовоз"
KIND_SURPLUS = "перетарка"
MARK_SHORT = "не-"
MARK_SURPLUS = "не+"

# Комментарий пишется в столбец N, заливка идёт по столбцам A:N.
COL_COMMENT = 14
SPAN_COLUMNS = 14

# Подписи столбцов реестра. Сравнение по началу строки без регистра.
HEAD_DAY = "дата"
HEAD_STORE = "магазин"
HEAD_DOC = "№ накладной"
HEAD_NAME = "номенклатура"
HEAD_SHORT = "кол-во недовоза"
HEAD_SURPLUS = "кол-во перетарки"

DATE_SHAPES = ("%d.%m.%Y", "%Y-%m-%d", "%d.%m.%y")

NO_BOOK = (
    "Реестр расхождений не загружен: недовозы и перетарки не размечены. "
    "Загрузите книгу на странице «Справочники»."
)
NO_DAY = (
    "Дата инвентаризации не определена из титула файла: без даты реестр "
    "не разбирается."
)
NO_SHEET = f"В книге реестра нет листа «{SHEET}»: заявки не прочитаны."


@dataclass
class Demand:
    """Заявка реестра: один товар одного вида по текущему магазину."""

    kind: str
    name: str
    key: str
    qty: float
    docs: tuple[str, ...] = ()

    @property
    def id(self) -> str:
        """Ключ подтверждения: не меняется между пересборками файла."""
        return md5(f"{self.kind}|{self.key}".encode("utf-8")).hexdigest()[:12]

    @property
    def mark(self) -> str:
        return MARK_SHORT if self.kind == KIND_SHORT else MARK_SURPLUS


@dataclass
class Mark:
    """Решение по одной заявке: какую строку сверки и как помечать."""

    key: str
    kind: str
    group: str
    row: int
    name: str
    claim_name: str
    docs: tuple[str, ...]
    qty: float
    demand_qty: float
    diff: float
    sum_before: float
    sum_after: float
    closed: bool
    mark: str
    comment: str

    def to_dict(self, applied: bool = False) -> dict:
        return {
            "key": self.key,
            "kind": self.kind,
            "group": self.group,
            "row": self.row,
            "name": self.name,
            "claim_name": self.claim_name,
            "docs": ", ".join(self.docs),
            "qty": self.qty,
            "demand_qty": self.demand_qty,
            "diff": self.diff,
            "sum_before": round(self.sum_before, 2),
            "sum_after": round(self.sum_after, 2),
            "closed": self.closed,
            "mark": self.mark,
            "comment": self.comment,
            "applied": applied,
        }


@dataclass
class Plan:
    """Что реестр предлагает сделать со строками сверки."""

    marks: list[Mark] = field(default_factory=list)
    skipped: list[str] = field(default_factory=list)


def book_path(refs_dir: str | Path) -> Path:
    return Path(refs_dir) / CLAIMS_FILE


def cache_path(refs_dir: str | Path) -> Path:
    return Path(refs_dir) / CACHE_FILE


def forget_cache(refs_dir: str | Path) -> None:
    """Убирает кеш разобранных строк: нужно после загрузки новой книги.

    Кеш и так проверяется по размеру и времени правки файла, но после
    ручной загрузки книги лишний файл на диске смысла не имеет.
    """
    try:
        cache_path(refs_dir).unlink(missing_ok=True)
    except OSError:
        logger.warning("Кеш реестра не удалён", exc_info=True)


def status(refs_dir: str | Path) -> dict:
    """Состояние книги реестра для страницы «Справочники»."""
    path = book_path(refs_dir)
    try:
        size = path.stat().st_size
    except OSError:
        return {"found": False, "size_mb": 0.0, "rows": 0}
    rows = 0
    try:
        rows = len(json.loads(cache_path(refs_dir).read_text("utf-8")).get("rows") or [])
    except Exception:  # noqa: BLE001
        rows = 0
    return {"found": True, "size_mb": round(size / 1024 / 1024, 2), "rows": rows}


def _text(value: object) -> str:
    return str(value if value is not None else "").strip()


def _day(value: object) -> dt.date | None:
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    text = _text(value)
    for shape in DATE_SHAPES:
        try:
            return dt.datetime.strptime(text, shape).date()
        except ValueError:
            continue
    return None


def _count(value: object) -> float:
    if isinstance(value, bool):
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = _text(value).replace(",", ".").replace(" ", "")
    try:
        return float(text)
    except ValueError:
        return 0.0


def _columns(header: list) -> dict[str, int]:
    """Номера нужных столбцов по подписям шапки."""
    wanted = {
        "day": HEAD_DAY,
        "store": HEAD_STORE,
        "doc": HEAD_DOC,
        "name": HEAD_NAME,
        "short": HEAD_SHORT,
        "surplus": HEAD_SURPLUS,
    }
    found: dict[str, int] = {}
    for number, value in enumerate(header, start=1):
        title = _text(value).lower().replace("ё", "е")
        if not title:
            continue
        for field_name, label in wanted.items():
            if field_name in found:
                continue
            if title.startswith(label.lower().replace("ё", "е")):
                found[field_name] = number
    return found


def read_book(path: str | Path) -> list[dict]:
    """Разбирает книгу реестра одним проходом.

    Строки без количества недовоза и перетарки пропускаются: в реестре
    много справочных записей, которые к сверке отношения не имеют.
    """
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if SHEET not in book.sheetnames:
            raise KeyError(NO_SHEET)
        sheet = book[SHEET]
        rows: list[dict] = []
        columns: dict[str, int] = {}
        for values in sheet.iter_rows(values_only=True):
            if not columns:
                columns = _columns(list(values))
                if len(columns) < 6:
                    columns = {}
                continue

            def at(name: str) -> object:
                number = columns.get(name, 0)
                return values[number - 1] if 0 < number <= len(values) else None

            short = _count(at("short"))
            surplus = _count(at("surplus"))
            if short <= 0 and surplus <= 0:
                continue
            day = _day(at("day"))
            name = _text(at("name"))
            store = _text(at("store"))
            if not name or not store or day is None:
                continue
            rows.append(
                {
                    "day": day.isoformat(),
                    "store": store,
                    "doc": _text(at("doc")),
                    "name": name,
                    "short": short,
                    "surplus": surplus,
                }
            )
        if not columns:
            raise KeyError(
                "В листе «" + SHEET + "» не найдена шапка: нужны столбцы "
                "«Дата», «Магазин», «№ накладной», «НОМЕНКЛАТУРА», "
                "«кол-во недовоза», «кол-во перетарки»."
            )
        return rows
    finally:
        book.close()


def _stamp(path: Path) -> list[float]:
    try:
        info = path.stat()
    except OSError:
        return [0.0, 0.0]
    return [float(info.st_size), float(info.st_mtime)]


_LOCK = threading.Lock()


def load_entries(refs_dir: str | Path, force: bool = False) -> list[dict]:
    """Строки реестра из кеша. Книга читается только при её смене."""
    path = book_path(refs_dir)
    if not path.is_file():
        return []
    stamp = _stamp(path)
    cache = cache_path(refs_dir)
    with _LOCK:
        if not force:
            try:
                saved = json.loads(cache.read_text("utf-8"))
                if saved.get("stamp") == stamp:
                    return list(saved.get("rows") or [])
            except Exception:  # noqa: BLE001
                pass
        rows = read_book(path)
        try:
            cache.parent.mkdir(parents=True, exist_ok=True)
            cache.write_text(
                json.dumps({"stamp": stamp, "rows": rows}, ensure_ascii=False),
                "utf-8",
            )
        except OSError:
            logger.warning("Кеш реестра не сохранён", exc_info=True)
        return rows


def store_entries(
    entries: list[dict],
    warehouse: str,
    day: dt.date,
    min_score: float,
) -> tuple[list[dict], list[str]]:
    """Строки текущего магазина за год строго назад от даты сверки.

    Имя магазина в реестре и в сверке пишут по-разному, поэтому сначала
    ищется точное совпадение после нормализации, а если его нет — самое
    похожее имя не ниже порога REFS_MATCH_MIN_SCORE.
    """
    notes: list[str] = []
    store = refs.normalize_store(warehouse)
    if not store:
        return [], [f"Имя склада «{warehouse}» после очистки пустое: реестр не разобран."]

    known = {refs.normalize_store(row["store"]) for row in entries}
    known.discard("")
    picked = store
    if store not in known:
        best, score = "", 0.0
        for name in sorted(known):
            ratio = refs.similarity(store, name)
            if ratio > score:
                best, score = name, ratio
        if not best or score < min_score:
            notes.append(
                f"Магазин «{warehouse}» в реестре расхождений не найден "
                f"(ближе всего «{best or "—"}», схожесть {score:.2f}, нужно "
                f"не меньше {min_score:.2f})."
            )
            return [], notes
        picked = best
        notes.append(
            f"Магазин «{warehouse}» сопоставлен с «{best}» в реестре, "
            f"схожесть {score:.2f}."
        )

    first = day - dt.timedelta(days=WINDOW_DAYS)
    picked_rows: list[dict] = []
    for row in entries:
        if refs.normalize_store(row["store"]) != picked:
            continue
        moment = _day(row["day"])
        if moment is None or not (first < moment <= day):
            continue
        picked_rows.append(row)
    if not picked_rows:
        notes.append(
            f"Заявок по магазину за год до {day:%d.%m.%Y} в реестре нет."
        )
    return picked_rows, notes


def comment(kind: str, docs: tuple[str, ...], qty: float) -> str:
    """Комментарий для столбца N: «недовоз ПДВ06030196 3шт.»."""
    numbers = " ".join(docs)
    pieces = f"{int(round(qty))}шт."
    return f"{kind} {numbers} {pieces}".replace("  ", " ").strip()


def demands(rows: list[dict], type_words: tuple[str, ...] | list[str]) -> list[Demand]:
    """Сводит строки реестра в заявки: одно название — одна заявка."""
    buckets: dict[tuple[str, str], Demand] = {}
    for row in rows:
        key = resort.normalize(row["name"], tuple(type_words))
        if not key:
            continue
        for kind, qty in ((KIND_SHORT, row["short"]), (KIND_SURPLUS, row["surplus"])):
            if qty <= 0:
                continue
            demand = buckets.get((kind, key))
            if demand is None:
                demand = Demand(kind=kind, name=row["name"], key=key, qty=0.0)
                buckets[(kind, key)] = demand
            demand.qty += float(qty)
            doc = row["doc"]
            if doc and doc not in demand.docs:
                demand.docs = demand.docs + (doc,)
    return sorted(buckets.values(), key=lambda item: (item.kind, item.name))


def _number(value: object) -> float:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return 0.0
    return float(value)


def load_demands(
    refs_dir: str | Path,
    warehouse: str,
    day: dt.date | None,
    min_score: float,
    type_words: tuple[str, ...] | list[str],
) -> tuple[list[Demand], list[str]]:
    """Заявки реестра по текущему магазину и дате инвентаризации."""
    if not book_path(refs_dir).is_file():
        return [], [NO_BOOK]
    if day is None:
        return [], [NO_DAY]
    try:
        entries = load_entries(refs_dir)
    except Exception as error:  # noqa: BLE001
        logger.exception("Реестр расхождений не прочитан")
        return [], [f"Реестр расхождений не прочитан: {type(error).__name__}: {error}"]
    rows, notes = store_entries(entries, warehouse, day, min_score)
    return demands(rows, type_words), notes


def sheet_rows(sheet, groups, type_words: tuple[str, ...] | list[str]) -> list[dict]:
    """Строки сверки с ключом названия: по ним ищутся заявки."""
    rows: list[dict] = []
    for group in groups:
        for row in group.data_rows:
            name = _text(sheet.cell(row=row, column=COL_NAME).value)
            if not name:
                continue
            rows.append(
                {
                    "group": group.name,
                    "row": row,
                    "name": name,
                    "key": resort.normalize(name, tuple(type_words)),
                    "diff": _number(sheet.cell(row=row, column=COL_DIFF).value),
                    "sum": _number(sheet.cell(row=row, column=COL_SUM_DIFF).value),
                }
            )
    return rows


def plan(rows: list[dict], wanted: list[Demand]) -> Plan:
    """Подбирает строки сверки под заявки реестра.

    Название должно совпасть один в один после нормализации, знак разницы
    должен отвечать виду заявки: недовоз — минус, перетарка — плюс.
    Подходить должна ровно одна строка: иначе это спорный случай и решает
    человек вручную.
    """
    result = Plan()
    for demand in wanted:
        candidates = [
            row
            for row in rows
            if row["key"] == demand.key
            and (row["diff"] < 0 if demand.kind == KIND_SHORT else row["diff"] > 0)
        ]
        if not candidates:
            result.skipped.append(
                f"{demand.kind} «{demand.name}» ({int(round(demand.qty))}шт.): "
                "в сверке нет подходящей строки."
            )
            continue
        if len(candidates) > 1:
            places = ", ".join(f"стр. {row['row']}" for row in candidates)
            result.skipped.append(
                f"{demand.kind} «{demand.name}»: спорный случай — подходит "
                f"несколько строк сверки ({places}). Разметьте вручную."
            )
            continue
        row = candidates[0]
        diff = row["diff"]
        size = abs(diff)
        qty = min(demand.qty, size)
        closed = demand.qty >= size
        if demand.qty > size:
            result.skipped.append(
                f"{demand.kind} «{demand.name}»: в заявках {int(round(demand.qty))}шт., "
                f"в сверке {int(round(size))}шт. — остаток заявки не размечен."
            )
        per_unit = row["sum"] / diff if diff else 0.0
        left = diff + qty if diff < 0 else diff - qty
        result.marks.append(
            Mark(
                key=demand.id,
                kind=demand.kind,
                group=row["group"],
                row=row["row"],
                name=row["name"],
                claim_name=demand.name,
                docs=demand.docs,
                qty=qty,
                demand_qty=demand.qty,
                diff=diff,
                sum_before=row["sum"],
                sum_after=0.0 if closed else round(per_unit * left, 2),
                closed=closed,
                mark=demand.mark,
                comment=comment(demand.kind, demand.docs, qty),
            )
        )
    return result


def confirmed(marks: list[Mark], decisions: dict[str, bool] | None) -> list[Mark]:
    """Заявки, подтверждённые человеком построчно."""
    answers = decisions or {}
    return [mark for mark in marks if answers.get(mark.key) is True]


def closed_rows(marks: list[Mark], decisions: dict[str, bool] | None) -> set[int]:
    """Строки, которые уходят из подбора пересортов.

    Из игры выходят только полностью закрытые реестром строки: остаток
    частично закрытой строки в пересортах участвует.
    """
    return {mark.row for mark in confirmed(marks, decisions) if mark.closed}


def apply(sheet, marks: list[Mark], moved: dict[int, int] | None = None) -> list[Mark]:
    """Пишет пометки, суммы, заливку и комментарии в файл сверки."""
    shift = moved or {}
    written: list[Mark] = []
    for mark in marks:
        row = shift.get(mark.row, mark.row)
        mark.row = row
        sum_cell = sheet.cell(row=row, column=COL_SUM_DIFF)
        if mark.closed:
            style.unmerge_at(sheet, row, COL_TRAIT)
            sheet.cell(row=row, column=COL_TRAIT).value = mark.mark
            sum_cell.value = 0
            for column in range(1, SPAN_COLUMNS + 1):
                style.paint(sheet.cell(row=row, column=column), style.GREEN)
        elif _number(sum_cell.value) != 0:
            # Пересорт мог обнулить сумму раньше: второй раз не вычитаем.
            sum_cell.value = mark.sum_after
        style.unmerge_at(sheet, row, COL_COMMENT)
        sheet.cell(row=row, column=COL_COMMENT).value = mark.comment
        written.append(mark)
    return written


def report(
    marks: list[Mark],
    decisions: dict[str, bool] | None,
    skipped: list[str],
    problems: list[str] | None = None,
) -> dict:
    """Итог реестра для страницы результата."""
    answers = decisions or {}
    rows = [mark.to_dict(applied=answers.get(mark.key) is True) for mark in marks]
    return {
        "rows": rows,
        "skipped": list(skipped),
        "problems": list(problems or []),
        "applied": sum(1 for row in rows if row["applied"]),
        "pending": sum(1 for row in rows if not row["applied"]),
    }
