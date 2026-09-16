"""Шаг 0б. Файл продавцов из 1С: ФИО и отработанные часы.

Выгрузка «Продавцы и часы» — отдельный файл на один склад:

* лист тот же, что у сверки (`TDSheet`);
* сверху блок параметров отбора: даты периода и строка
  `Склад Равно "ИмяСклада"`;
* затем заголовки `Продавец` и `Часы`;
* затем строки парами: ФИО, под ним число часов;
* последняя строка — итог часов, своего ФИО у него нет.

Строки объединены по `A:C`, поэтому обычно читается столбец `A`; если
выгрузка положит часы в соседний столбец, они тоже будут найдены.

Модуль ничего не пишет в сверку: он отдаёт готовый список `Seller`,
который веб-слой кладёт в `SheetMeta.sellers` — дальше блок продавцов
рисует `meta.write_sellers()` так же, как при ручном вводе.
"""

from __future__ import annotations

import logging
import re
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

from .guard import check_archive
from .meta import Seller, hours_text
from .repair import repair

logger = logging.getLogger("excelkro.sellers")

# Заголовки и служебные строки выгрузки: в список продавцов не попадают.
HEADER_WORDS = ("продавец", "часы", "сотрудник", "итого", "всего")
SERVICE_WORDS = ("параметры", "отбор", "дата нач", "дата кон", "валюта", "склад")
WAREHOUSE_PATTERN = re.compile(r"склад[^\"«]*[\"«]([^\"»]+)[\"»]", re.IGNORECASE)
# Столбцы A:C — ФИО и часы объединённой строки; D — на случай другой выгрузки.
MAX_COLUMN = 4
MAX_ROWS = 2000


class SellersError(Exception):
    """Файл продавцов не удалось прочитать."""


@dataclass
class SellersFile:
    """Разобранный файл продавцов."""

    sellers: tuple[Seller, ...] = field(default_factory=tuple)
    warehouse: str = ""
    total: float | int | None = None
    notes: list[str] = field(default_factory=list)

    @property
    def hours_sum(self) -> float:
        return sum(float(item.hours) for item in self.sellers if item.hours)

    def note(self) -> str:
        """Короткая строка для страницы хода разбора."""
        if not self.sellers:
            return "продавцы в файле не найдены"
        hours = self.hours_sum
        plain = int(hours) if hours == int(hours) else round(hours, 2)
        return f"продавцов: {len(self.sellers)}, часов: {plain}"


def _text(value: object) -> str:
    return str(value if value is not None else "").strip()


def _number(value: object) -> float | None:
    """Число из ячейки: принимаем запятую, пробелы и текстовую запись."""
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    plain = _text(value).replace(",", ".").replace(" ", "").replace("\xa0", "")
    if not plain:
        return None
    try:
        return float(plain)
    except ValueError:
        return None


def _is_service(text: str) -> bool:
    low = text.lower()
    return any(low.startswith(word) for word in SERVICE_WORDS)


def _is_header(text: str) -> bool:
    return text.lower().strip(": ") in HEADER_WORDS


def warehouse_from_filter(text: str) -> str:
    """Имя склада из строки отбора «Склад Равно \"ИмяСклада\"»."""
    found = WAREHOUSE_PATTERN.search(text or "")
    return found.group(1).strip() if found else ""


def _row_values(sheet, row: int) -> list[object]:
    return [
        sheet.cell(row=row, column=column).value
        for column in range(1, MAX_COLUMN + 1)
    ]


def read_sheet(sheet) -> SellersFile:
    """Список продавцов с листа выгрузки.

    Правило простое и не зависит от номеров строк: текст — это ФИО,
    число — часы предыдущего ФИО. Число без ФИО перед ним считается
    итогом и в список не попадает.
    """
    warehouse = ""
    result: list[Seller] = []
    total: float | int | None = None
    notes: list[str] = []
    limit = min(int(sheet.max_row or 0), MAX_ROWS)

    for row in range(1, limit + 1):
        values = _row_values(sheet, row)
        texts = [_text(value) for value in values]
        line = " ".join(item for item in texts if item)
        if not line:
            continue
        if not warehouse:
            warehouse = warehouse_from_filter(line)

        name = ""
        hours: float | None = None
        for value, text in zip(values, texts):
            if not text:
                continue
            number = _number(value)
            if number is not None:
                if hours is None:
                    hours = number
                continue
            if _is_header(text) or _is_service(text):
                continue
            if not name:
                name = text

        if name:
            result.append(
                Seller(name=name, hours=hours_text(hours) if hours is not None else "")
            )
            continue
        if hours is not None:
            if result and not result[-1].hours:
                result[-1] = Seller(name=result[-1].name, hours=hours_text(hours))
            else:
                # Число без ФИО и без пустых часов — итог внизу выгрузки.
                total = int(hours) if hours == int(hours) else hours

    empty = [item.name for item in result if not item.hours]
    if empty:
        notes.append(
            "Часы не распознаны у продавцов: "
            + ", ".join(empty)
            + ". Впишите их вручную на странице результата."
        )
    return SellersFile(sellers=tuple(result), warehouse=warehouse, total=total, notes=notes)


def read(
    path: str | Path,
    folder: str | Path,
    sheet_name: str = "TDSheet",
    repair_mode: str = "inject",
    max_unpacked_mb: int = 200,
) -> SellersFile:
    """Читает файл продавцов из 1С. Ремонт архива обязателен, как у сверки."""
    path = Path(path)
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    try:
        check_archive(path, max_unpacked_mb)
        repaired = repair(path, folder / "sellers_repaired.xlsx", repair_mode)
        workbook = openpyxl.load_workbook(repaired, data_only=True)
    except Exception as error:  # noqa: BLE001
        raise SellersError(
            f"Файл продавцов не прочитан: {type(error).__name__}: {error}"
        ) from error
    sheet = (
        workbook[sheet_name]
        if sheet_name in workbook.sheetnames
        else workbook.worksheets[0]
    )
    data = read_sheet(sheet)
    logger.info("Продавцов из файла: %s", len(data.sellers))
    return data
