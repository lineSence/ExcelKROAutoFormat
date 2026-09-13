"""Разбор книги графика ревизий: кто и где делал переучет.

Лист «Переучеты» устроен так:

* в первом столбце — дата, по одной строке на день;
* шапка «дата | администраторы…» повторяется каждую неделю, и состав
  столбцов со временем меняется, причём администраторов пишут то
  фамилией («Казаков А.»), то просто именем («Георгий», «Света»);
* в ячейке на пересечении даты и администратора — магазин и под ним
  фамилии ревизоров через перевод строки.

Разметка плавает, поэтому угадывать фамилию по виду слова нельзя: названия
магазинов («Жемчужина», «Думская», «Ульянка») выглядят точно так же, как
фамилии. Поэтому ревизоры узнаются по списку людей из самой книги
(листы с телефонами), а не по шаблону слова:

1. собирается список ФИО с листов «тлф», «Тлф помощников», «Подотчетники»;
2. строка внутри ячейки считается ревизором, если её фамилия есть в списке
   (с запасом на опечатки: «Палова М.» -> «Павлова»);
3. всё остальное — название магазина или пометка о работе («сбор товара»,
   «открытие», «+1 продавца»).

Магазин в ячейке стоит не всегда первым и не всегда есть вовсе. Его пишут
один раз, а дальше в ячейках ниже остаются только фамилии — так ведут
многодневные переучёты и ночные инвентаризации, где сверху стоит магазин,
а снизу — люди. Поэтому название продолжается вниз по столбцу, а найденные
фамилии относятся и к своей дате, и к дате той ячейки, где был магазин.

Важно о чтении: книга открывается в режиме `read_only`, где обращение
`sheet.cell(...)` каждый раз заново разбирает весь XML листа, поэтому лист
читается одним проходом в матрицу значений.
"""

from __future__ import annotations

import datetime as dt
import logging
import re
from difflib import SequenceMatcher
from pathlib import Path

import openpyxl

logger = logging.getLogger("excelkro.schedule")

SCHEDULE_SHEET = "Переучеты"

# Листы книги графика, где лежат полные ФИО и телефоны.
NAME_SHEETS = ("тлф", "тлф помощников", "график помощников", "подотчетники")

# Страховка от гигантских листов.
MAX_ROWS = 20000
MAX_COLUMNS = 200

# Полное ФИО: три слова с большой буквы.
FULL_NAME = re.compile(r"^([А-ЯЁ][а-яё\-]+)\s+([А-ЯЁ][а-яё]+)\s+([А-ЯЁ][а-яё]+)$")
# Разделители фамилий внутри одной строки ячейки.
PARTS_SPLIT = re.compile(r"[;,/+\u2022\u00b7|]|\sи\s")

# Слова служебных пометок: ни фамилия, ни название магазина.
SERVICE_WORDS = (
    "сбор",
    "товар",
    "открыт",
    "закрыт",
    "выклад",
    "перестанов",
    "переуч",
    "инвентариз",
    "ревизи",
    "ревизор",
    "помощник",
    "продав",
    "отдел",
    "ночь",
    "дежур",
    "отпуск",
    "больнич",
    "выходн",
    "дата",
    "итог",
    "склад",
    "магазин",
)

# Насколько похожа фамилия из графика на фамилию из списка людей.
NAME_MIN_SCORE = 0.88
# Короче этого слово фамилией не считается.
MIN_SURNAME = 4
# Сколько столбцов с именами делают строку шапкой администраторов.
MIN_ADMINS = 3


def _clean(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _key(value: str) -> str:
    return value.lower().replace("ё", "е")


def _name_case(text: str) -> str:
    """«ИВАНОВА М.» -> «Иванова М.»: в графике фамилии часто пишут капсом."""
    words = []
    for word in _clean(text).split():
        if word.isalpha() and word.isupper() and len(word) > 2:
            word = word.capitalize()
        words.append(word)
    return " ".join(words)


def _is_service(text: str) -> bool:
    low = _key(text)
    return any(word in low for word in SERVICE_WORDS)


def normalize_store(value: object) -> str:
    """Имя склада без уточнений в скобках, регистра и знаков."""
    text = str(value or "").lower().replace("ё", "е")
    text = re.sub(r"\([^)]*\)", " ", text)
    return re.sub(r"[^0-9a-zа-я]+", "", text)


def _as_date(value: object) -> dt.date | None:
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    text = _clean(value)
    for shape in ("%d.%m.%Y", "%d.%m.%y", "%Y-%m-%d"):
        try:
            return dt.datetime.strptime(text, shape).date()
        except ValueError:
            continue
    return None


def _matrix(sheet) -> list[list]:
    """Значения листа списком строк. Один проход по файлу."""
    rows: list[list] = []
    for number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
        rows.append(list(values[:MAX_COLUMNS]))
        if number >= MAX_ROWS:
            logger.warning("Лист «%s» обрезан на %s строках", getattr(sheet, "title", "?"), MAX_ROWS)
            break
    return rows


def _width(rows: list[list]) -> int:
    return max((len(line) for line in rows), default=0)


def _at(rows: list[list], row: int, column: int) -> object:
    if 1 <= row <= len(rows):
        line = rows[row - 1]
        if 1 <= column <= len(line):
            return line[column - 1]
    return None


# --- Список людей -------------------------------------------------------


def read_people(book) -> dict[str, list[str]]:
    """Фамилия -> полные ФИО с листов с телефонами."""
    by_surname: dict[str, list[str]] = {}
    for name in book.sheetnames:
        if _key(name.strip()) not in NAME_SHEETS:
            continue
        for row in book[name].iter_rows(values_only=True):
            for value in row:
                for part in str(value or "").split("\n"):
                    text = _name_case(part)
                    match = FULL_NAME.match(text)
                    if not match:
                        continue
                    names = by_surname.setdefault(_key(match.group(1)), [])
                    if text not in names:
                        names.append(text)
    return by_surname


def surname_in(text: str, people: dict[str, list[str]]) -> tuple[str, float]:
    """Фамилия из списка, которой отвечает кусок ячейки.

    Пустая строка в ответе — это не человек из списка.
    """
    text = _name_case(text)
    if not text or any(char.isdigit() for char in text) or _is_service(text):
        return "", 0.0
    first = re.sub(r"[^а-яё\-]+", "", _key(text).split()[0])
    if len(first) < MIN_SURNAME:
        return "", 0.0
    if first in people:
        return first, 1.0
    best, score = "", 0.0
    for candidate in people:
        ratio = SequenceMatcher(None, first, candidate).ratio()
        if ratio > score:
            best, score = candidate, ratio
    if score < NAME_MIN_SCORE:
        return "", score
    return best, score


def _cell_content(
    raw: object, people: dict[str, list[str]]
) -> tuple[list[str], list[str]]:
    """Разбор одной ячейки: фамилии ревизоров и остальные строки.

    Строка целиком считается названием магазина, если хотя бы один её
    кусок не из списка людей: «Бухарестская, 47» не должна распадаться
    на два значения по запятой.
    """
    found: list[str] = []
    others: list[str] = []
    for line in str(raw or "").split("\n"):
        text = _clean(line)
        if not text:
            continue
        parts = [piece for piece in (_clean(item) for item in PARTS_SPLIT.split(text)) if piece]
        named = [piece for piece in parts if surname_in(piece, people)[0]]
        for piece in named:
            name = _name_case(piece)
            if name not in found:
                found.append(name)
        if len(named) != len(parts):
            others.append(text)
    return found, others


# --- Шапка с администраторами ----------------------------------------------


def _admin_columns(rows: list[list], row: int) -> dict[int, str]:
    """Столбцы администраторов в строке шапки.

    В шапке стоит то фамилия («Казаков А.»), то одно имя («Георгий»),
    поэтому список людей здесь не помогает: берём короткие слова с
    большой буквы без цифр и служебных слов.
    """
    found: dict[int, str] = {}
    for column in range(2, _width(rows) + 1):
        text = _clean(_at(rows, row, column))
        if not text or any(char.isdigit() for char in text) or _is_service(text):
            continue
        if len(text.split()) > 2 or not text[0].isupper():
            continue
        found[column] = _name_case(text)
    return found


def _is_header(rows: list[list], row: int) -> bool:
    """Строка «дата | администраторы…»: повторяется по всему листу."""
    first = _key(_clean(_at(rows, row, 1)))
    return first.startswith("дат") and len(_admin_columns(rows, row)) >= MIN_ADMINS


# --- Разбор листа графика -------------------------------------------------


def _remember(
    visits: dict[tuple[str, dt.date], tuple[str, tuple[str, ...]]],
    store: str,
    day: dt.date,
    admin: str,
    found: list[str],
) -> None:
    """Добавляет запись графика, собирая людей из всех её ячеек.

    Один магазин в один день встречается у двух администраторов и в
    соседних строках (ночная инвентаризация), поэтому фамилии копятся.
    """
    known = visits.get((store, day))
    if known is None:
        visits[(store, day)] = (admin, tuple(found))
        return
    merged = list(known[1])
    for name in found:
        if name not in merged:
            merged.append(name)
    visits[(store, day)] = (known[0] or admin, tuple(merged))


def read_schedule(path: str | Path) -> tuple[
    dict[tuple[str, dt.date], tuple[str, tuple[str, ...]]],
    dict[str, list[str]],
]:
    """Читает книгу графика: распределение по датам и список ФИО."""
    visits: dict[tuple[str, dt.date], tuple[str, tuple[str, ...]]] = {}

    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        people = read_people(book)
        if SCHEDULE_SHEET not in book.sheetnames:
            logger.warning("В книге графика нет листа %s", SCHEDULE_SHEET)
            return visits, people
        rows = _matrix(book[SCHEDULE_SHEET])
    finally:
        book.close()

    if not people:
        logger.warning(
            "В книге графика нет листов со списком людей (%s): ревизоры не распознаются",
            ", ".join(NAME_SHEETS),
        )

    admins: dict[int, str] = {}
    # Магазин пишут только в первый день: при многодневном переучёте и при
    # ночной инвентаризации, когда сверху название, а снизу фамилии.
    carried: dict[int, tuple[str, dt.date]] = {}
    for row in range(1, len(rows) + 1):
        day = _as_date(_at(rows, row, 1))
        if day is None:
            if _is_header(rows, row):
                admins = _admin_columns(rows, row)
                carried = {}
            continue
        if not admins:
            continue
        for column, admin in admins.items():
            raw = _at(rows, row, column)
            if not _clean(raw):
                continue
            found, others = _cell_content(raw, people)
            title = ""
            for text in others:
                if _is_service(text):
                    continue
                title = text
                break
            if title:
                origin = day
            else:
                # Фамилии без названия: магазин взят из ячейки выше по столбцу.
                title, origin = carried.get(column, ("", day))
            store = normalize_store(title)
            if not store:
                continue
            carried[column] = (title, origin)
            _remember(visits, store, day, admin, found)
            if origin != day and found:
                # Ночная инвентаризация: фамилии относятся и к дате магазина.
                _remember(visits, store, origin, admin, found)
    return visits, people


def debug_schedule_block(path: str | Path, day: dt.date, limit: int = 200) -> list[str]:
    """Сырые ячейки строки графика на указанную дату.

    Нужна для разбора «неудобных» книг: показывает, что реально лежит
    в ячейках и какие куски признаны фамилиями.
    """
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        people = read_people(book)
        if SCHEDULE_SHEET not in book.sheetnames:
            return [f"В книге графика нет листа {SCHEDULE_SHEET}"]
        rows = _matrix(book[SCHEDULE_SHEET])
    finally:
        book.close()

    out: list[str] = [f"Фамилий в списке людей: {len(people)}"]
    for row in range(1, len(rows) + 1):
        if _as_date(_at(rows, row, 1)) != day:
            continue
        out.append(f"Дата найдена в строке {row}")
        for column in range(1, _width(rows) + 1):
            raw = _at(rows, row, column)
            if not _clean(raw):
                continue
            found, others = _cell_content(raw, people)
            out.append(f"  c{column} {raw!r}")
            out.append(f"      ревизоры: {found} | остальное: {others}")
            if len(out) >= limit:
                out.append("  … вывод обрезан")
                return out
    if len(out) == 1:
        out.append("Дата в первом столбце листа не найдена")
    return out
