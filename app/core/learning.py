"""Режим обучения: датасет из ручных образцов и обучение модели.

Источники примеров:
1. Готовые ручные сверки (файл загружается в разделе «Обучение»).
   Пара в рамке (medium) — положительный пример, прочие допустимые
   пары внутри группы — отрицательные. Зелёные строки (ручной тег `не-`)
   в обучение не попадают.
2. Решения пользователя по спорным парам: каждое «пересорт» / «не пересорт»
   сразу становится разметкой. Так модель улучшается со временем.

Храним не признаки, а сырые поля пары: признаки можно поменять и пересчитать
датасет заново. Формат хранилища — JSONL, одна строка на пример.
"""

from __future__ import annotations

import json
import logging
import random
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from pathlib import Path

from .resort import (
    Item,
    make_item,
    price_allowed,
    same_brand,
    similarity,
)
from .verify import FEATURE_NAMES, LogisticModel, pair_features, train_logistic

logger = logging.getLogger("excelkro.learning")

COL_NAME = 2
COL_DIFF = 6
COL_PRICE = 7
COL_SUM_DIFF = 10
GREEN = "FF92D050"
CANDIDATE_SIMILARITY = 0.5


@dataclass
class Sample:
    """Один пример для обучения: пара строк и ответ человека."""

    group: str
    plus_name: str
    plus_diff: float
    plus_sum: float
    plus_row: int
    minus_name: str
    minus_diff: float
    minus_sum: float
    minus_row: int
    label: int
    source: str
    origin: str = "sample"
    added: str = ""

    def key(self) -> str:
        return "|".join(
            [
                self.source,
                self.group,
                self.plus_name.strip().lower(),
                self.minus_name.strip().lower(),
                f"{self.plus_diff:g}",
                f"{self.minus_diff:g}",
            ]
        )

    def items(self, type_words: tuple[str, ...]) -> tuple[Item, Item]:
        plus = make_item(self.plus_row, self.plus_name, self.plus_diff, self.plus_sum, type_words)
        minus = make_item(self.minus_row, self.minus_name, self.minus_diff, self.minus_sum, type_words)
        return plus, minus


# ——— чтение ручного образца ———


def _fill_code(cell) -> str:
    fill = cell.fill
    if fill is None or fill.fill_type != "solid":
        return ""
    value = fill.start_color.rgb
    return value if isinstance(value, str) else ""


def _framed(cell) -> bool:
    border = cell.border
    return any(
        getattr(border, side) and getattr(border, side).style == "medium"
        for side in ("left", "right", "top", "bottom")
    )


def _number(value: object) -> float:
    return float(value) if isinstance(value, (int, float)) and not isinstance(value, bool) else 0.0


def _group_bounds(sheet) -> list[tuple[int, int, str]]:
    """Границы групп готового файла: по формуле итога в столбце J."""
    bounds: list[tuple[int, int, str]] = []
    for row in range(1, sheet.max_row + 1):
        value = sheet.cell(row=row, column=COL_SUM_DIFF).value
        if not isinstance(value, str) or not value.startswith("=SUM(J"):
            continue
        inside = value[value.index("(") + 1 : value.index(")")]
        try:
            first = int(inside.split(":")[0][1:])
        except ValueError:
            continue
        name = ""
        for probe in range(first - 1, max(0, first - 4), -1):
            text = str(sheet.cell(row=probe, column=COL_NAME).value or "").strip()
            if text and text != "№":
                name = text
                break
        bounds.append((first, row - 1, name or f"Группа {first}"))
    return bounds


def samples_from_manual(
    path: str | Path,
    type_words: tuple[str, ...],
    source: str | None = None,
    max_negatives_per_group: int = 400,
) -> list[Sample]:
    """Разметка из ручного образца: рамка = пара, остальное = не пара."""
    import openpyxl

    file = Path(path)
    name = source or file.name
    book = openpyxl.load_workbook(file)
    sheet = book[book.sheetnames[0]]
    stamp = datetime.now(timezone.utc).isoformat(timespec="seconds")

    samples: list[Sample] = []
    for first, last, group_name in _group_bounds(sheet):
        items: list[Item] = []
        blocks: dict[int, int] = {}
        block_number = 0
        previous_framed = False

        for row in range(first, last + 1):
            diff_cell = sheet.cell(row=row, column=COL_DIFF)
            diff = _number(diff_cell.value)
            framed = _framed(diff_cell)
            if framed and not previous_framed:
                block_number += 1
            previous_framed = framed
            if diff == 0:
                continue
            if _fill_code(diff_cell) == GREEN:
                # Ручной тег по внешним данным: не учимся на нём.
                continue
            price = _number(sheet.cell(row=row, column=COL_PRICE).value)
            sum_diff = _number(sheet.cell(row=row, column=COL_SUM_DIFF).value)
            if sum_diff == 0:
                # В готовом файле сумма пересорта обнулена: считаем её из цены.
                sum_diff = diff * price
            item = make_item(row, sheet.cell(row=row, column=COL_NAME).value or "", diff, sum_diff, type_words)
            items.append(item)
            if framed:
                blocks[row] = block_number

        pluses = [item for item in items if item.diff > 0]
        minuses = [item for item in items if item.diff < 0]
        negatives: list[Sample] = []
        for plus in pluses:
            for minus in minuses:
                block_plus = blocks.get(plus.row)
                block_minus = blocks.get(minus.row)
                linked = block_plus is not None and block_plus == block_minus
                if not linked:
                    plausible = (
                        same_brand(plus, minus, 0.80)
                        or price_allowed(plus, minus, 0.80)
                        or similarity(plus.key, minus.key) >= CANDIDATE_SIMILARITY
                    )
                    if not plausible:
                        continue
                sample = Sample(
                    group=group_name,
                    plus_name=plus.name,
                    plus_diff=plus.diff,
                    plus_sum=plus.sum_diff,
                    plus_row=plus.row,
                    minus_name=minus.name,
                    minus_diff=minus.diff,
                    minus_sum=minus.sum_diff,
                    minus_row=minus.row,
                    label=1 if linked else 0,
                    source=name,
                    origin="sample",
                    added=stamp,
                )
                if linked:
                    samples.append(sample)
                else:
                    negatives.append(sample)

        if len(negatives) > max_negatives_per_group:
            # Берём самые похожие отрицательные примеры: они самые полезные.
            negatives.sort(
                key=lambda row: similarity(
                    make_item(0, row.plus_name, 1, 1, type_words).key,
                    make_item(0, row.minus_name, 1, 1, type_words).key,
                ),
                reverse=True,
            )
            negatives = negatives[:max_negatives_per_group]
        samples.extend(negatives)

    book.close()
    return samples


# ——— хранилище ———


def load_samples(path: str | Path) -> list[Sample]:
    file = Path(path)
    if not file.is_file():
        return []
    rows: list[Sample] = []
    for line in file.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line:
            continue
        try:
            rows.append(Sample(**json.loads(line)))
        except (ValueError, TypeError):
            continue
    return rows


def append_samples(path: str | Path, samples: list[Sample], limit: int = 60000) -> dict:
    """Добавляет примеры без дублей. Возвращает сводку."""
    file = Path(path)
    existing = load_samples(file)
    known = {row.key() for row in existing}
    fresh = [row for row in samples if row.key() not in known]
    for row in fresh:
        known.add(row.key())

    merged = existing + fresh
    if len(merged) > limit:
        merged = merged[len(merged) - limit :]

    file.parent.mkdir(parents=True, exist_ok=True)
    with file.open("w", encoding="utf-8") as handle:
        for row in merged:
            handle.write(json.dumps(asdict(row), ensure_ascii=False) + "\n")

    return {"added": len(fresh), "skipped": len(samples) - len(fresh), "total": len(merged)}


def clear_samples(path: str | Path) -> None:
    file = Path(path)
    if file.is_file():
        file.unlink()


def dataset_stats(path: str | Path) -> dict:
    rows = load_samples(path)
    sources: dict[str, int] = {}
    for row in rows:
        sources[row.source] = sources.get(row.source, 0) + 1
    return {
        "total": len(rows),
        "positives": sum(1 for row in rows if row.label),
        "negatives": sum(1 for row in rows if not row.label),
        "from_decisions": sum(1 for row in rows if row.origin == "decision"),
        "sources": sorted(sources.items()),
    }


# ——— обучение ———


def _metrics(model: LogisticModel, vectors: list[list[float]], labels: list[int]) -> dict:
    if not vectors:
        return {}
    true_positive = false_positive = false_negative = correct = 0
    for vector, label in zip(vectors, labels):
        predicted = 1 if model.predict(vector) >= 0.5 else 0
        correct += predicted == label
        true_positive += predicted == 1 and label == 1
        false_positive += predicted == 1 and label == 0
        false_negative += predicted == 0 and label == 1
    precision = true_positive / (true_positive + false_positive) if true_positive + false_positive else 0.0
    recall = true_positive / (true_positive + false_negative) if true_positive + false_negative else 0.0
    f1 = 2 * precision * recall / (precision + recall) if precision + recall else 0.0
    return {
        "pairs": len(vectors),
        "accuracy": round(correct / len(vectors), 4),
        "precision": round(precision, 4),
        "recall": round(recall, 4),
        "f1": round(f1, 4),
    }


def train_model(settings) -> dict:
    """Обучает модель на всём хранилище и сохраняет её в JSON."""
    rows = load_samples(settings.train_store_path)
    if len(rows) < 40:
        raise ValueError("Для обучения нужно не меньше 40 примеров.")
    if not any(row.label for row in rows) or not any(not row.label for row in rows):
        raise ValueError("Нужны и пары, и не-пары.")

    type_words = tuple(settings.type_words)
    vectors: list[list[float]] = []
    labels: list[int] = []
    for row in rows:
        plus, minus = row.items(type_words)
        vectors.append(pair_features(plus, minus))
        labels.append(int(row.label))

    # Проверка на отложенных файлах: делим по источникам, а не по парам.
    sources = sorted({row.source for row in rows})
    holdout_metrics: dict = {}
    if len(sources) >= 2:
        random.Random(17).shuffle(sources)
        holdout = set(sources[: max(1, len(sources) // 5)])
        train_index = [index for index, row in enumerate(rows) if row.source not in holdout]
        test_index = [index for index, row in enumerate(rows) if row.source in holdout]
        if train_index and test_index:
            probe = train_logistic(
                [vectors[index] for index in train_index],
                [labels[index] for index in train_index],
                epochs=int(settings.train_epochs),
            )
            holdout_metrics = _metrics(
                probe,
                [vectors[index] for index in test_index],
                [labels[index] for index in test_index],
            )
            holdout_metrics["files"] = sorted(holdout)

    model = train_logistic(vectors, labels, epochs=int(settings.train_epochs))
    model.meta = {
        "trained_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "samples": len(rows),
        "positives": sum(labels),
        "sources": sorted({row.source for row in rows}),
        "epochs": int(settings.train_epochs),
        "train_metrics": _metrics(model, vectors, labels),
        "holdout_metrics": holdout_metrics,
        "features": list(FEATURE_NAMES),
    }
    model.save(settings.model_path)
    logger.info("Модель обучена на %s примерах", len(rows))
    return model.meta


def samples_from_decisions(
    rows: list[dict],
    answers: dict[str, bool],
    source: str,
) -> list[Sample]:
    """Примеры из решений пользователя на странице результата."""
    stamp = datetime.now(timezone.utc).isoformat(timespec="seconds")
    out: list[Sample] = []
    for row in rows:
        key = str(row.get("key") or "")
        if key not in answers:
            continue
        payload = row.get("pair") or {}
        if not payload.get("plus_name") or not payload.get("minus_name"):
            continue
        out.append(
            Sample(
                group=str(row.get("group") or ""),
                plus_name=str(payload["plus_name"]),
                plus_diff=float(payload.get("plus_diff") or 0),
                plus_sum=float(payload.get("plus_sum") or 0),
                plus_row=int(payload.get("plus_row") or 0),
                minus_name=str(payload["minus_name"]),
                minus_diff=float(payload.get("minus_diff") or 0),
                minus_sum=float(payload.get("minus_sum") or 0),
                minus_row=int(payload.get("minus_row") or 0),
                label=1 if answers[key] else 0,
                source=source,
                origin="decision",
                added=stamp,
            )
        )
    return out
