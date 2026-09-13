"""Цвета, рамки, ширины и высоты готового файла."""

from __future__ import annotations

from copy import copy

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.colors import Color
from openpyxl.utils.cell import range_boundaries

YELLOW = "FFFFFF00"
ORANGE = "FFFFC000"
GREEN = "FF92D050"
BLACK = "FF000000"

# В образце формат без разделителя тысяч.
MONEY_FORMAT = "0.00"
DATE_FORMAT = "DD.MM.YYYY"

# В образце коробка общего итога шириной в две колонки: I:J.
TOTAL_SPAN = 2
# Коробка значения неучтёнки тоже шириной в две колонки: C:D.
EXTRA_SPAN = 2
# Плашка срока приёма найденного товара занимает E:J.
NOTICE_SPAN = 6
# Заливка плашки в образце задана темой книги, а не цветом RGB.
NOTICE_THEME = 5

# Ширины столбцов 1–12. Остальные остаются по умолчанию.
COLUMN_WIDTHS = {
    "A": 5.0,
    "B": 55.0,
    "C": 5.0,
    "D": 5.0,
    "E": 5.0,
    "F": 5.0,
    "G": 10.0,
    "H": 10.0,
    "I": 10.0,
    "J": 10.0,
    "K": 55.0,
    "L": 55.0,
}
ROW_HEIGHT = 12

# Excel показывает ширину на 0,83 меньше записанного значения:
# в файле ширина хранится вместе с отступами ячейки.
# Поэтому к нужной ширине добавляется эта поправка.
WIDTH_PADDING = 0.83

# Свойства оформления, которые переносятся при перестановке строк.
# Используются только открытые свойства openpyxl: скрытое `cell._style`
# меняется от версии к версии.
STYLE_FIELDS = (
    "font",
    "fill",
    "border",
    "alignment",
    "protection",
    "number_format",
)


def read_style(cell) -> dict:
    """Снимок оформления ячейки."""
    return {name: copy(getattr(cell, name)) for name in STYLE_FIELDS}


def apply_style(cell, saved: dict) -> None:
    """Ставит ранее снятое оформление на ячейку."""
    for name in STYLE_FIELDS:
        if name in saved:
            setattr(cell, name, copy(saved[name]))


def copy_style(source, target) -> None:
    """Переносит оформление с одной ячейки на другую."""
    apply_style(target, read_style(source))


def fill(color: str) -> PatternFill:
    return PatternFill(fill_type="solid", start_color=color, end_color=color)


def theme_fill(theme: int = NOTICE_THEME) -> PatternFill:
    """Заливка цветом темы книги, как у плашки в образце."""
    return PatternFill(fill_type="solid", fgColor=Color(theme=theme, tint=0.0))


def no_fill() -> PatternFill:
    """Пустая заливка."""
    return PatternFill(fill_type=None)


def no_border() -> Border:
    """Отсутствие рамки."""
    return Border()


def medium_side() -> Side:
    return Side(style="medium", color=BLACK)


def medium_border() -> Border:
    side = medium_side()
    return Border(left=side, right=side, top=side, bottom=side)


def paint(cell, color: str) -> None:
    cell.fill = fill(color)


def clear(cell) -> None:
    """Снимает заливку и рамку: ячейка остаётся чистой."""
    cell.fill = no_fill()
    cell.border = no_border()


def frame(cell) -> None:
    cell.border = medium_border()


def unmerge_at(sheet, row: int, column: int) -> None:
    """Снимает объединение, если ячейка в него входит."""
    for text in [str(item) for item in sheet.merged_cells.ranges]:
        min_col, min_row, max_col, max_row = range_boundaries(text)
        if min_row <= row <= max_row and min_col <= column <= max_col:
            sheet.unmerge_cells(text)


def wide_box(cell, span: int = TOTAL_SPAN) -> None:
    """Растягивает коробку итога на две колонки, как в образце (I:J).

    Оформление берётся с левой ячейки: openpyxl сам раскидывает
    рамку по контуру объединённого диапазона.
    """
    sheet = getattr(cell, "parent", None)
    if sheet is None or span < 2:
        return
    row = cell.row
    column = cell.column
    for shift in range(span):
        unmerge_at(sheet, row, column + shift)
    sheet.merge_cells(
        start_row=row,
        start_column=column,
        end_row=row,
        end_column=column + span - 1,
    )


def frame_block(sheet, rows, column: int) -> None:
    """Ставит толстую рамку по контуру блока строк.

    Соседние строки собираются в один блок. Внутри блока
    горизонтальных линий нет, как в образце.
    """
    ordered = sorted(set(int(row) for row in rows))
    if not ordered:
        return
    side = medium_side()
    runs: list[list[int]] = []
    for row in ordered:
        if runs and row == runs[-1][-1] + 1:
            runs[-1].append(row)
        else:
            runs.append([row])
    for run in runs:
        for row in run:
            cell = sheet.cell(row=row, column=column)
            cell.border = Border(
                left=side,
                right=side,
                top=side if row == run[0] else None,
                bottom=side if row == run[-1] else None,
            )


def style_total_cell(cell) -> None:
    """Итог группы в столбце J."""
    cell.font = Font(name="Arial", size=9, bold=True)
    cell.number_format = MONEY_FORMAT
    cell.border = medium_border()
    cell.alignment = Alignment(horizontal="center", vertical="center")


def style_grand_total_cell(cell) -> None:
    """Общий итог в столбце I.

    В образце итог стоит в жёлтой коробке со средней рамкой,
    шириной в две колонки (I:J), число по центру.
    """
    cell.font = Font(name="Arial", size=10, bold=True)
    cell.number_format = MONEY_FORMAT
    cell.border = medium_border()
    cell.alignment = Alignment(horizontal="center", vertical="center")
    paint(cell, YELLOW)
    wide_box(cell)


def style_net_total_cell(cell) -> None:
    """Итог с неучтёнкой (строка «С неучтёнкой:»).

    В образце это зелёная коробка I:J со средней рамкой:
    именно от этого числа считается ставка продавцов.
    """
    cell.font = Font(name="Arial", size=10, bold=True)
    cell.number_format = MONEY_FORMAT
    cell.border = medium_border()
    cell.alignment = Alignment(horizontal="center", vertical="center")
    paint(cell, GREEN)
    wide_box(cell)


def style_extra_label_cell(cell) -> None:
    """Подписи блока неучтёнки в B1 и B2.

    Жёлтая заливка, средняя рамка, Arial 10 полужирный, текст влево.
    """
    cell.font = Font(name="Arial", size=10, bold=True)
    cell.number_format = "General"
    cell.border = medium_border()
    cell.alignment = Alignment(horizontal="left", vertical="center")
    paint(cell, YELLOW)


def style_extra_value_cell(cell) -> None:
    """Значения неучтёнки в C1 и C2.

    В образце это жёлтая коробка на две колонки (C:D) со средней
    рамкой и числом по центру.
    """
    cell.font = Font(name="Arial", size=10, bold=True)
    cell.number_format = "General"
    cell.border = medium_border()
    cell.alignment = Alignment(horizontal="center", vertical="center")
    paint(cell, YELLOW)
    wide_box(cell, EXTRA_SPAN)


def style_notice_cell(cell) -> None:
    """Плашка «Найденный товар принимается до:» в E1:J1.

    В образце это широкая ячейка с заливкой темы, текстом
    по центру и средней линией слева — она отделяет плашку
    от коробки неучтёнки.
    """
    cell.font = Font(name="Arial", size=10)
    cell.number_format = "General"
    cell.border = Border(left=medium_side())
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.fill = theme_fill()
    wide_box(cell, NOTICE_SPAN)


def style_field_label_cell(cell, right_line: bool = False, bottom_line: bool = False) -> None:
    """Подпись поля шапки: «Склад:», «Недостача:» и прочие.

    В образце все подписи Arial 9 и прижаты вправо, к своему
    значению. Линии по краю нужны только у блока итогов.
    """
    cell.font = Font(name="Arial", size=9)
    cell.alignment = Alignment(horizontal="right", vertical="center")
    if right_line or bottom_line:
        side = medium_side()
        cell.border = Border(
            right=side if right_line else None,
            bottom=side if bottom_line else None,
        )


def style_field_value_cell(cell, span: int = 1, bottom_line: bool = False) -> None:
    """Значение поля шапки: имя склада, причина инвентаризации."""
    cell.font = Font(name="Arial", size=9, bold=True)
    cell.alignment = Alignment(horizontal="left", vertical="center")
    if bottom_line:
        cell.border = Border(bottom=medium_side())
    if span > 1:
        wide_box(cell, span)


def apply_geometry(sheet, last_row: int) -> None:
    """Шаг 8: ширины, высоты, автофильтр, параметры печати."""
    for letter, width in COLUMN_WIDTHS.items():
        sheet.column_dimensions[letter].width = width + WIDTH_PADDING
    for row in range(1, max(last_row, sheet.max_row) + 1):
        sheet.row_dimensions[row].height = ROW_HEIGHT
    sheet.auto_filter.ref = f"K1:K{last_row}"
    sheet.page_setup.orientation = "portrait"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
