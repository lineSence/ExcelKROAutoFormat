"""Связка всех шагов: ремонт → разбор → формат → справочники → сохранение."""

from __future__ import annotations

import datetime as dt
import logging
import shutil
import time
import uuid
from dataclasses import asdict, dataclass, field
from pathlib import Path

import openpyxl

from ..config import Settings
from . import prev as prev_book
from . import refs, report
from .format import format_workbook, output_filename
from .guard import check_archive
from .meta import SheetMeta
from .parse import ParseError, warehouse_from_filename
from .refs_sync import local_paths
from .repair import RepairError, repair

logger = logging.getLogger("excelkro.pipeline")

DATE_SHAPES = ("%d.%m.%Y", "%d.%m.%y", "%Y-%m-%d")

NO_PREV_DATE = (
    "Дата предыдущей инвентаризации не определена из титула загруженного файла: "
    "поле осталось как было — впишите дату вручную."
)


@dataclass
class PipelineResult:
    """Результат обработки одного файла."""

    output_path: Path
    output_name: str
    summary: dict
    groups: list[dict]
    clusters: list[dict]
    doubtful: list[dict]
    source_path: Path | None = None
    source_name: str = ""
    decisions: dict[str, bool] | None = None
    # Причина, администратор, проверяющий и ревизоры из справочников.
    refs: dict = field(default_factory=dict)
    # Ручные поля сверки: нужны при каждой пересборке файла.
    sheet_meta: SheetMeta = field(default_factory=SheetMeta)
    # Сравнение с предыдущей инвентаризацией и её файл.
    comparison: dict = field(default_factory=dict)
    prev_path: Path | None = None
    prev_name: str = ""


def work_dir(settings: Settings) -> Path:
    folder = Path(settings.tmp_dir) / uuid.uuid4().hex
    folder.mkdir(parents=True, exist_ok=True)
    return folder


def sweep(settings: Settings) -> int:
    """Удаляет рабочие папки старше срока хранения. Возвращает число папок."""
    root = Path(settings.tmp_dir)
    if not root.is_dir():
        return 0
    deadline = time.time() - max(settings.result_ttl_minutes, 1) * 60
    removed = 0
    for folder in root.iterdir():
        try:
            if not folder.is_dir() or folder.stat().st_mtime >= deadline:
                continue
        except OSError:
            continue
        shutil.rmtree(folder, ignore_errors=True)
        removed += 1
    if removed:
        logger.info("Удалено старых рабочих папок: %s", removed)
    return removed


def pick_sheet(workbook, settings: Settings):
    if settings.sheet_name in workbook.sheetnames:
        return workbook[settings.sheet_name]
    if not workbook.sheetnames:
        raise ParseError("В файле нет ни одного листа.")
    logger.warning(
        "Лист %s не найден. Взят первый лист: %s",
        settings.sheet_name,
        workbook.sheetnames[0],
    )
    return workbook[workbook.sheetnames[0]]


def load_sheet(repaired: Path, settings: Settings):
    """Открывает книгу и даёт понятное сообщение при битом файле."""
    try:
        workbook = openpyxl.load_workbook(repaired)
    except (IndexError, KeyError, ValueError) as error:
        raise RepairError(
            "В файле биты ссылки на стили или таблица текстов. "
            "Пересохраните выгрузку в Excel или включите запасной ремонт: "
            "REPAIR_MODE=libreoffice."
        ) from error
    return workbook, pick_sheet(workbook, settings)


def doc_day(value: object) -> dt.date | None:
    """Дата инвентаризации из титула в виде даты."""
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    text = str(value or "").strip()
    for shape in DATE_SHAPES:
        try:
            return dt.datetime.strptime(text, shape).date()
        except ValueError:
            continue
    return None


def prev_date_text(previous) -> str:
    """Дата из титула загруженного файла предыдущей инвентаризации.

    Возвращает вид `ГГГГ-ММ-ДД`: так дата сразу показывается в поле формы
    и понимается шагом записи в файл.
    """
    day = doc_day(getattr(previous, "date", "") or "")
    return day.isoformat() if day else ""


def apply_prev_date(
    sheet_meta: SheetMeta | None,
    previous,
    notes: list[str],
) -> SheetMeta | None:
    """Ставит дату предыдущей инвентаризации по загруженному файлу.

    Дата из самого документа точнее ручного ввода, поэтому она заменяет
    значение в ручных полях и показывается в форме на странице результата.
    Если дата в титуле не разобрана, ручное значение остаётся нетронутым.
    """
    if previous is None:
        return sheet_meta
    text = prev_date_text(previous)
    if not text:
        notes.append(NO_PREV_DATE)
        return sheet_meta
    meta = sheet_meta if sheet_meta is not None else SheetMeta()
    meta.prev_date = text
    logger.info("Дата предыдущей инвентаризации взята из файла: %s", text)
    return meta


def _closest_store(store: str, names: list[str]) -> tuple[str, float]:
    """Ближайшее известное имя склада и оценка схожести."""
    best, score = "", 0.0
    for name in names:
        ratio = refs.similarity(store, name)
        if ratio > score:
            best, score = name, ratio
    return best, score


def _visit_days(books: refs.RefsBooks, store: str) -> list[dt.date]:
    return sorted({day for name, day in books.visits if name == store})


def refs_problems(
    warehouse: str,
    day: dt.date | None,
    books: refs.RefsBooks,
    settings: Settings,
) -> list[str]:
    """Готовит понятные причины, почему справочники не дали данных.

    Сообщения показываются на странице результата и на странице «Справочники».
    В сам файл сверки они не попадают никогда.
    """
    planning, schedule = local_paths(settings.refs_dir)
    problems: list[str] = []

    if not planning.is_file():
        problems.append(
            "Книга планирования не загружена: причину инвентаризации брать неоткуда."
        )
    elif not books.plans:
        problems.append(
            "Книга планирования прочитана, но ни одного магазина не распознано. "
            "Ожидается: на листе администратора во второй строке недели вида «Январь 10-16», "
            "во втором столбце названия магазинов, данные с третьей строки."
        )

    if not schedule.is_file():
        problems.append(
            "Книга графика не загружена: администратора и ревизоров брать неоткуда."
        )
    elif not books.visits:
        problems.append(
            f"В книге графика не разобрано распределение (лист «{refs.SCHEDULE_SHEET}»). "
            "Ожидается: даты в первом столбце, шапка с фамилиями администраторов "
            "(не меньше трёх в строке), в ячейке — магазин, а под ним фамилии ревизоров."
        )

    if day is None:
        problems.append(
            "Дата инвентаризации не определена из титула файла: без даты справочники "
            "не ищутся. Проверьте ячейку с датой в выгрузке 1С."
        )

    store = refs.normalize_store(warehouse)
    if not store:
        problems.append(
            f"Имя склада «{warehouse}» после очистки пустое: проверьте имя файла сверки."
        )
        return problems

    threshold = settings.refs_match_min_score

    if books.visits:
        names = sorted({name for name, _ in books.visits})
        best, score = _closest_store(store, names)
        if score < threshold:
            problems.append(
                f"Склад «{warehouse}» не найден в графике (администратор и ревизоры). "
                f"Ближе всего «{best}», схожесть {score:.2f}, нужно не меньше {threshold:.2f} "
                "(настройка REFS_MATCH_MIN_SCORE)."
            )
        elif day is not None:
            days = _visit_days(books, best)
            near = [
                one
                for one in days
                if abs((one - day).days) <= settings.refs_days_around
            ]
            if not near:
                span = f"{days[0]:%d.%m.%Y} — {days[-1]:%d.%m.%Y}" if days else "записей нет"
                problems.append(
                    f"Склад в графике есть («{best}»), но на {day:%d.%m.%Y} "
                    f"±{settings.refs_days_around} дн. записи нет. Даты этого склада в графике: {span}."
                )

    if books.plans:
        names = sorted(books.plans)
        best, score = _closest_store(store, names)
        if score < threshold:
            problems.append(
                f"Склад «{warehouse}» не найден в книге планирования (причина). "
                f"Ближе всего «{best}», схожесть {score:.2f}, нужно не меньше {threshold:.2f}."
            )
        elif day is not None:
            weeks = books.plans.get(best, [])
            if not any(start <= day <= end for start, end, _, _ in weeks):
                problems.append(
                    f"Склад в планировании есть («{best}»), но неделя с датой {day:%d.%m.%Y} "
                    "не заполнена или не распознана: причина осталась пустой."
                )

    return problems


def fill_refs(sheet, warehouse: str, day: dt.date | None, settings: Settings) -> dict:
    """Подставляет данные справочников в готовый файл.

    Точные совпадения пишутся сразу. Если имя склада совпало неточно или
    запись взята со сдвигом даты, значения считаются неуверенными: в файл
    они не попадают, а показываются на странице результата на подтверждение.
    Проверяющий всегда один и тот же, поэтому он пишется всегда.
    Служебные пометки в сверку не пишутся никогда.
    """
    planning, schedule = local_paths(settings.refs_dir)
    books = refs.load_books(str(planning), str(schedule))
    info = refs.lookup(
        warehouse,
        day,
        books,
        checker=settings.default_checker,
        min_score=settings.refs_match_min_score,
        days_around=settings.refs_days_around,
        confirm_min_score=settings.refs_confirm_min_score,
    )
    # Неуверенные значения ждут подтверждения человека.
    to_write = refs.pending(info) if info.uncertain else info
    written = refs.write_cells(sheet, to_write, settings.refs_cells())
    # Разбор причин — тяжёлый шаг, поэтому считается ровно один раз.
    need_problems = not info.found or not info.reason
    problems = refs_problems(warehouse, day, books, settings) if need_problems else []
    return {
        "reason": info.reason,
        "admin": info.admin,
        "checker": info.checker,
        "auditors": list(info.auditors),
        "found": info.found,
        "cells": written,
        "problems": problems,
        # Подтверждение неуверенных значений.
        "uncertain": info.uncertain,
        "confidence": round(info.confidence, 2),
        "day_shift": info.day_shift,
        "notes": list(info.notes),
        "counts": {
            "plans": len(books.plans),
            "visits": len(books.visits),
            "people": len(books.by_surname),
        },
    }


def read_previous(
    prev_path: str | Path,
    prev_name: str,
    folder: Path,
    settings: Settings,
) -> tuple[object | None, list[str]]:
    """Готовит предыдущую инвентаризацию к сравнению.

    Сбой второго файла не должен мешать обычной обработке сверки.
    """
    try:
        check_archive(prev_path, settings.max_unpacked_mb)
        repaired = repair(prev_path, folder / "prev_repaired.xlsx", settings.repair_mode)
        previous = prev_book.read(
            repaired,
            prev_name or Path(prev_path).name,
            settings.sheet_name,
        )
    except Exception as error:  # noqa: BLE001
        logger.exception("Предыдущая сверка не прочитана")
        return None, [
            f"Предыдущая сверка не прочитана: {type(error).__name__}: {error}"
        ]
    if not previous.items:
        return previous, [
            "В предыдущей сверке не найдено ни одной позиции: сравнивать нечего."
        ]
    return previous, []


def process(
    input_path: str | Path,
    original_filename: str,
    settings: Settings | None = None,
    decisions: dict[str, bool] | None = None,
    folder: str | Path | None = None,
    sheet_meta: SheetMeta | None = None,
    prev_path: str | Path | None = None,
    prev_name: str = "",
) -> PipelineResult:
    """Обрабатывает файл сверки и возвращает путь к готовому файлу.

    `folder` — готовая рабочая папка. Веб-слой передаёт ту же папку, в которую
    сохранил загруженный файл, чтобы лишние папки не оставались на диске.
    `sheet_meta` — ручные поля сверки (причина, продавцы, подписи).
    `prev_path` — файл предыдущей инвентаризации для сравнения (необязательно).
    Его дата из титула становится датой предыдущей инвентаризации в сверке.
    """
    settings = settings or Settings.load()
    folder = Path(folder) if folder is not None else work_dir(settings)
    folder.mkdir(parents=True, exist_ok=True)

    check_archive(input_path, settings.max_unpacked_mb)
    repaired = repair(input_path, folder / "repaired.xlsx", settings.repair_mode)

    workbook, sheet = load_sheet(repaired, settings)

    warehouse = warehouse_from_filename(original_filename)
    if not warehouse:
        raise ParseError("Из имени файла не вышло получить имя склада.")

    previous = None
    prev_problems: list[str] = []
    if prev_path:
        previous, prev_problems = read_previous(prev_path, prev_name, folder, settings)
        # Дата предыдущей инвентаризации берётся из самого файла.
        sheet_meta = apply_prev_date(sheet_meta, previous, prev_problems)

    format_result = format_workbook(
        sheet,
        warehouse,
        settings,
        decisions,
        sheet_meta,
        previous,
    )

    comparison: dict = {}
    if prev_path:
        comparison = asdict(format_result.comparison)
        comparison["credit_sum"] = round(float(comparison.get("credit_sum") or 0.0), 2)
        comparison["file_name"] = comparison.get("file_name") or prev_name
        comparison["items"] = len(previous.items) if previous is not None else 0
        comparison["sellers"] = len(previous.sellers) if previous is not None else 0
        comparison["problems"] = prev_problems
        comparison["prev_date"] = sheet_meta.prev_date if sheet_meta is not None else ""

    try:
        refs_info = fill_refs(
            sheet,
            warehouse,
            doc_day(format_result.document.doc_date),
            settings,
        )
    except Exception as error:  # noqa: BLE001
        # Сбой справочников не должен мешать выдаче файла сверки.
        logger.exception("Справочники не применены")
        refs_info = {
            "reason": "",
            "admin": "",
            "checker": settings.default_checker,
            "auditors": [],
            "found": False,
            "cells": [],
            "problems": [
                f"Сбой при работе со справочниками: {type(error).__name__}: {error}"
            ],
            "uncertain": False,
            "confidence": 0.0,
            "day_shift": 0,
            "notes": [],
            "counts": {},
        }

    output_name = output_filename(warehouse, format_result.document.doc_date)
    output_path = folder / output_name
    workbook.save(output_path)

    return PipelineResult(
        output_path=output_path,
        output_name=output_name,
        summary=report.summary(format_result),
        groups=report.group_rows(format_result),
        clusters=report.cluster_rows(format_result) if settings.report_cluster_members else [],
        doubtful=report.doubtful_rows(format_result),
        source_path=Path(input_path),
        source_name=original_filename,
        decisions=dict(decisions or {}),
        refs=refs_info,
        sheet_meta=sheet_meta or SheetMeta(),
        comparison=comparison,
        prev_path=Path(prev_path) if prev_path else None,
        prev_name=prev_name,
    )
